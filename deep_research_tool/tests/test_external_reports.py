"""External research material in the feedback re-research (起案 C):
reading md / docx / html / pasted text (chapters, cited URLs, digests),
the LLM comparison and its rule-based fallback, the comparison turned
into planner instructions, the material as secondary-information
documents, revise_report end to end (comparison → plan → cited URLs
fetched as primary sources → external chapters cited with
access_method "external" and capped importance → comparison table in
the result) and the Web API (upload paths, pasted text saved to a file,
precheck). No network, no paid LLM."""
import json
from pathlib import Path

import pytest

from deep_research_tool import external_reports as ext
from deep_research_tool.tests.test_concurrency_visibility import _post, web  # noqa: F401
from deep_research_tool.tests.test_e2e_pipeline import FakeSearchClient
from deep_research_tool.tests.test_finalization_integration import build_tool
from deep_research_tool.tests.test_revision import (FEEDBACK, RevLLM, TestWebApi as _RevWeb,
                                                    _run_v1, _session)

ODR_MD = """# 炭素繊維市場の調査（OpenAI Deep Research）

## 1. 市場規模

国内の炭素繊維市場は2025年に2,400億円へ拡大した[1]。前年比8%の成長である。

## 2. 主要メーカーの戦略

東レは約1,000億円の増産投資を発表した[2]。帝人も追随している。

## 3. 中国メーカーの台頭

中国メーカーの生産能力は2025年に世界の約4割に達した（[吉林化繊の公表資料](https://jilin.example/ir)）。

## Sources
[1] 経済産業省 炭素繊維統計 2025: https://meti.example/cf2025
[2] 東レ IR: https://toray.example/ir/2025
"""

COMPARE_JSON = json.dumps({"topics": [
    {"topic": "市場規模", "status": "conflict", "section": "1", "report": "1,900 億円（2023 年）",
     "external": "2,400 億円（2025 年）", "label": "ODR", "note": "年が異なる"},
    {"topic": "中国メーカーの台頭", "status": "external_only", "section": None, "report": "",
     "external": "世界の約4割", "label": "ODR"},
    {"topic": "主要メーカーの戦略", "status": "agree", "section": "2", "report": "増産投資",
     "external": "増産投資", "label": "ODR"},
    {"topic": "不明", "status": "weird"},
    {"topic": "航空機需要", "status": "report_only", "section": "9", "report": "回復"}]}, ensure_ascii=False)


class ExtLLM(RevLLM):
    """RevLLM plus an answer for the comparison prompt; records prompts."""
    compare_json = COMPARE_JSON

    def __init__(self):
        super().__init__()
        self.prompts = []

    def _answer(self, prompt):
        self.prompts.append(prompt)
        if "外部の調査資料（章と要旨・数値）" in prompt:
            return self.compare_json
        return super()._answer(prompt)


# ----------------------------------------------------------------------
# reading
# ----------------------------------------------------------------------

class TestLoading:
    def test_markdown_paste_and_cited_urls(self, tmp_path):
        f = tmp_path / "odr.md"
        f.write_text(ODR_MD, encoding="utf-8")
        reps = ext.load_external_reports([{"label": "ODR", "path": str(f)},
                                          {"text": "## 概要\n\n貼り付けた本文。https://paste.example/a を参照。"}])
        a, b = reps
        assert a.error is None and a.origin == "file" and a.label == "ODR"
        assert a.title.startswith("炭素繊維市場の調査")
        assert [s["title"] for s in a.sections] == ["市場規模", "主要メーカーの戦略", "中国メーカーの台頭"]
        # the Sources section is not a chapter, but its URLs are cited (with titles) first
        assert [u["url"] for u in a.urls] == ["https://meti.example/cf2025", "https://toray.example/ir/2025",
                                              "https://jilin.example/ir"]
        assert a.urls[0]["title"] == "経済産業省 炭素繊維統計 2025" and a.urls[2]["title"] == "吉林化繊の公表資料"
        assert "2,400億円" in a.sections[0]["digest"] and "（数値:" not in a.sections[0]["digest"]
        assert b.origin == "paste" and b.label == "外部資料2" and [u["url"] for u in b.urls] == ["https://paste.example/a"]
        info = a.info()
        assert info["file"] == "odr.md" and info["sections"] == 3 and info["cited_urls"] == 3

    def test_docx_html_and_errors(self, tmp_path):
        import docx
        d = docx.Document()
        d.add_heading("Gemini の調査", level=0)
        d.add_heading("市場", level=1)
        d.add_paragraph("市場は拡大している。詳細は https://gem.example/x を参照。")
        d.save(tmp_path / "g.docx")
        (tmp_path / "p.html").write_text(
            "<html><head><title>Copilot</title></head><body><h2>動向</h2><p>本文。</p>"
            "<a href='https://cop.example/1'>出典</a></body></html>", encoding="utf-8")
        (tmp_path / "x.xlsx").write_bytes(b"PK")
        (tmp_path / "empty.md").write_text("   ", encoding="utf-8")
        reps = ext.load_external_reports([
            {"path": str(tmp_path / "g.docx")}, {"path": str(tmp_path / "p.html")},
            {"path": str(tmp_path / "x.xlsx")}, {"path": str(tmp_path / "empty.md")},
            {"path": str(tmp_path / "missing.md")}])
        g, h, x, e, m = reps
        assert g.error is None and g.sections and [u["url"] for u in g.urls] == ["https://gem.example/x"]
        assert h.error is None and [u["url"] for u in h.urls] == ["https://cop.example/1"]
        assert "対応していない形式" in x.error and "空" in e.error and "見つかりません" in m.error
        with pytest.raises(ValueError, match="5 件まで"):
            ext.load_external_reports([{"text": "a"}] * 6)
        # truncation and label de-dup
        reps = ext.load_external_reports([{"label": "A", "text": "あ" * 50}, {"label": "A", "text": "い" * 50}],
                                         max_chars=20)
        assert [r.label for r in reps] == ["A", "A (2)"] and all(r.truncated and r.chars == 20 for r in reps)

    def test_urls_end_at_japanese_punctuation_and_keep_balanced_parens(self):
        text = ("（出典：https://www.meti.go.jp/report/a.pdf）によると拡大。【https://b.example/x】"
                "see https://en.wikipedia.org/wiki/Foo_(bar). and [w](https://en.wikipedia.org/wiki/Baz_(qux)) "
                "https://ja.wikipedia.org/wiki/炭素繊維、です「https://c.example/y」")
        assert [u["url"] for u in ext.extract_cited_urls(text)] == [
            "https://en.wikipedia.org/wiki/Baz_(qux)", "https://www.meti.go.jp/report/a.pdf",
            "https://b.example/x", "https://en.wikipedia.org/wiki/Foo_(bar)",
            "https://ja.wikipedia.org/wiki/炭素繊維", "https://c.example/y"]
        # a URL removed from a digest takes nothing after it
        assert ext.digest("国内市場（出典：https://meti.example/cf）は2025年に2,400億円へ拡大した。") == \
            "国内市場（出典：）は2025年に2,400億円へ拡大した。"

    def test_html_headings_and_unclosed_links(self):
        title, text = ext.html_to_markdown(
            "<title>T</title><h2>\n  Market size\n</h2><p>a <a href='https://x.example/1'>link</a></p>"
            "<h2>B<br>part</h2><p>body <a href='https://y.example/2'>open<h2>C</h2><p>body c</p>"
            "<script>var x = 1;</script><p>tail <a href='https://z.example/3'>never closed")
        assert title == "T"
        assert text == ("## Market size\n\na [link](https://x.example/1)\n\n## B part\n\n"
                        "body [open](https://y.example/2)\n\n## C\n\nbody c\n\ntail [never closed](https://z.example/3)")

    def test_helpers(self):
        urls = ext.extract_cited_urls("a https://x.example/p?utm_source=z and https://x.example/p) "
                                      "[t](https://y.example/q)", limit=5)
        assert [u["url"] for u in urls] == ["https://y.example/q", "https://x.example/p?utm_source=z"]
        assert len(ext.extract_cited_urls(" ".join(f"https://h{i}.example/" for i in range(50)))) == ext.MAX_CITED_URLS
        dg = ext.digest("# 見出し\n\n" + "長い文。" * 100 + "市場は2025年に3兆円、シェアは40%。")
        assert dg.startswith("見出し 長い文。") and dg.endswith("（数値: 2025年、3兆円、40%）")


# ----------------------------------------------------------------------
# comparison
# ----------------------------------------------------------------------

class TestComparison:
    def _reports(self, tmp_path):
        f = tmp_path / "odr.md"
        f.write_text(ODR_MD, encoding="utf-8")
        return ext.load_external_reports([{"label": "ODR", "path": str(f)}])

    def test_llm_comparison_and_instructions(self, tmp_path):
        s = _session()
        reps = self._reports(tmp_path)
        llm = ExtLLM()
        comp = ext.compare_reports(llm, s, reps)
        assert comp.source == "llm"
        assert [(t.topic, t.status, t.section) for t in comp.topics] == [
            ("市場規模", "conflict", "1"), ("中国メーカーの台頭", "external_only", None),
            ("主要メーカーの戦略", "agree", "2"), ("航空機需要", "report_only", None)]
        prompt = llm.prompts[0]
        assert "■ ODR" in prompt and "市場規模の推移" in prompt and "2,400億円" in prompt
        text = ext.comparison_instructions(comp, "add")
        assert "第1章の「市場規模」: 外部資料（ODR）は「2,400 億円（2025 年）」" in text
        assert "一次資料で確認し、違いがあれば両論を根拠つきで書く" in text
        assert "新しい章として追加する" in text and "記録のみ・章は変えない）: 主要メーカーの戦略" in text
        assert "章は追加しない" in ext.comparison_instructions(comp, "revise")
        assert ext.comparison_instructions(ext.Comparison(), "add") == ""
        assert "【食い違い】市場規模 第1章" in ext.comparison_preview(comp)

    def test_fallback_and_none(self, tmp_path):
        s = _session()
        reps = self._reports(tmp_path)
        broken = ExtLLM()
        broken.compare_json = "nope"
        comp = ext.compare_reports(broken, s, reps)
        assert comp.source == "fallback" and len(comp.topics) == 3
        assert {t.status for t in comp.topics} <= {"agree", "external_only"}
        assert all(t.label == "ODR" and t.external for t in comp.topics)
        assert ext.compare_reports(broken, s, []).source == "none"
        with pytest.raises(ValueError):
            ext.parse_comparison('{"topics": [{"topic": "x", "status": "??"}]}', s, reps)

    def test_documents_and_summary(self, tmp_path):
        reps = self._reports(tmp_path)
        docs = ext.external_documents(reps)
        assert len(docs) == 3 and docs[0]["title"] == "外部調査（ODR）: 市場規模"
        assert docs[0]["metadata"]["source"] == "external" and docs[0]["path"] == "external/ODR/01"
        assert ext.cited_urls(reps, limit=2) == ["https://meti.example/cf2025", "https://toray.example/ir/2025"]
        comp = ext.compare_reports(ExtLLM(), _session(), reps)

        class Ev:
            def __init__(self, sec, method):
                self.section_reference, self.access_method, self.metadata = sec, method, {}

        class Locker:
            def get_all_evidence(self):
                return [Ev("1", "web"), Ev("1", "external"), Ev("3", "external")]
        rev_ = {"revised": ["1"], "added": ["3"]}
        url_report = {"urls": [{"url": "https://meti.example/cf2025", "ok": True},
                               {"url": "https://toray.example/ir/2025", "ok": False}]}
        out = ext.external_summary(reps, comp, locker=Locker(), revision=rev_,
                                   cited=["https://meti.example/cf2025", "https://toray.example/ir/2025"],
                                   url_report=url_report)
        rows = {t["topic"]: t for t in out["topics"]}
        assert rows["市場規模"]["outcome"] == "一次資料 1 件で再調査"
        assert rows["主要メーカーの戦略"]["outcome"] == "一致（そのまま）"
        assert out["counts"] == {"agree": 1, "conflict": 1, "external_only": 1, "report_only": 1}
        assert out["cited_fetched"] == 1 and out["cited_failed"] == ["https://toray.example/ir/2025"]
        assert out["external_only_sections"] == ["3"]
        assert rows["中国メーカーの台頭"]["outcome"] == "対応する章なし（修正計画で扱われず）"
        # the chapter the plan added for an external-only topic is found by its title
        rev2 = {"revised": ["1"], "added": ["3"],
                "sections": [{"section": "1", "title": "市場規模の推移", "action": "revise"},
                             {"section": "3", "title": "中国メーカーの台頭と影響", "action": "add"}]}
        out = ext.external_summary(reps, comp, locker=Locker(), revision=rev2)
        row = {t["topic"]: t for t in out["topics"]}["中国メーカーの台頭"]
        assert row["matched_section"] == "3"
        assert row["outcome"] == "第3章を追加・一次資料なし（外部資料のみ・要確認）"
        # a shared word is not a match ("メーカーの" in both titles)
        rev3 = {"revised": ["2"], "added": [],
                "sections": [{"section": "2", "title": "主要メーカーの動向", "action": "revise"}]}
        out = ext.external_summary(reps, comp, revision=rev3)
        row = {t["topic"]: t for t in out["topics"]}["中国メーカーの台頭"]
        assert "matched_section" not in row and row["outcome"].startswith("対応する章なし")

    def test_summary_counts_only_this_runs_evidence_and_full_mode(self, tmp_path):
        reps = self._reports(tmp_path)
        comp = ext.compare_reports(ExtLLM(), _session(), reps)

        class Ev:
            def __init__(self, id_, sec, method):
                self.id, self.section_reference, self.access_method, self.metadata = id_, sec, method, {}

        class Locker:
            def get_all_evidence(self):
                return [Ev("old1", "1", "web"), Ev("old2", "1", "web"), Ev("new1", "1", "external")]
        rev_ = {"revised": ["1"], "added": [], "sections": [{"section": "1", "title": "市場規模の推移", "action": "revise"}]}
        out = ext.external_summary(reps, comp, locker=Locker(), revision=rev_)
        assert {t["topic"]: t for t in out["topics"]}["市場規模"]["outcome"] == "一次資料 2 件で再調査"
        # the original report's evidence (copied into the new locker) does not count
        out = ext.external_summary(reps, comp, locker=Locker(), revision=rev_,
                                   parent_evidence_ids={"old1", "old2"})
        assert {t["topic"]: t for t in out["topics"]}["市場規模"]["outcome"] == "一次資料なし（外部資料のみ・要確認）"
        assert out["external_only_sections"] == ["1"]
        # full mode: the new report's own chapters, matched by title (never by number)
        full = {"mode": "full", "revised": [], "added": [],
                "sections": [{"section": "1", "title": "中国メーカーの台頭", "action": "keep"},
                             {"section": "2", "title": "市場規模の推移", "action": "keep"}]}

        class FullLocker:
            def get_all_evidence(self):
                return [Ev("a", "1", "web"), Ev("b", "2", "external")]
        out = ext.external_summary(reps, comp, locker=FullLocker(), revision=full)
        rows = {t["topic"]: t for t in out["topics"]}
        assert rows["中国メーカーの台頭"]["matched_section"] == "1"
        assert rows["中国メーカーの台頭"]["outcome"] == "全面再調査・第1章・一次資料 1 件で再調査"
        assert rows["市場規模"]["matched_section"] == "2"
        assert rows["市場規模"]["outcome"] == "全面再調査・第2章・一次資料なし（外部資料のみ・要確認）"
        assert out["external_only_sections"] == ["2"]


# ----------------------------------------------------------------------
# revise_report end to end
# ----------------------------------------------------------------------

class TestToolWithExternalMaterial:
    def test_add_mode_with_material_only(self, tmp_path):
        tool1, res1 = _run_v1(tmp_path)
        parent_path = res1["saved_artifacts"]["session"]
        f = tmp_path / "odr.md"
        f.write_text(ODR_MD, encoding="utf-8")
        llm = ExtLLM()
        search = FakeSearchClient()
        tool2 = build_tool(tmp_path, llm, search, enable_verification=False, research_iterations=1)
        tool2.config.research.use_enhanced_synthesis = False
        tool2.config.research.max_gap_fill_rounds = 0
        shown = []

        def review(plan, revise_fn):
            shown.append(plan)
            return plan
        msgs = []
        res = tool2.revise_report(parent_session_path=parent_path, feedback="", mode="add",
                                  external_reports=[{"label": "ODR", "path": str(f)},
                                                    {"label": "壊れ", "path": str(tmp_path / "none.md")}],
                                  plan_review_callback=review,
                                  progress_callback=lambda m, p: msgs.append(m))
        r = res["revision"]
        assert res["kind"] == "revision" and r["revised"] == ["2"] and r["added"] == ["3"]
        x = r["external"]
        assert [m["label"] for m in x["materials"]] == ["ODR", "壊れ"] and x["materials"][1]["error"]
        assert x["comparison_source"] == "llm" and x["counts"]["conflict"] == 1
        assert x["cited_urls"] == ["https://meti.example/cf2025", "https://toray.example/ir/2025",
                                   "https://jilin.example/ir"] and x["cited_fetched"] == 3
        assert x["body_as_evidence"] is True
        # the planner got the comparison as feedback; the review shows the table
        plan_prompt = next(p for p in llm.prompts if "修正計画" in p and "JSON のみを出力" in p)
        assert "【外部の調査資料との照合結果】" in plan_prompt and "2,400 億円（2025 年）" in plan_prompt
        assert "外部資料との照合:" in shown[0].summary and "【食い違い】市場規模" in shown[0].summary
        assert any("照合" in m for m in msgs) and any("引用する URL 3 件" in m for m in msgs)
        # the writing rule joined the requirements; cited pages are primary (url), chapters secondary
        s2 = tool2.researcher.get_session()
        assert "【出典の書き方】" in s2.requirements and "外部調査（出所名）によれば" in s2.requirements
        assert s2.run_config["external_reports"][0]["label"] == "ODR"
        evs = tool2.researcher.get_evidence_locker().get_all_evidence()
        ext_evs = [e for e in evs if e.access_method == "external"]
        assert ext_evs and all(e.title.startswith("外部調査（ODR）") for e in ext_evs)
        assert all((e.importance_score or 0) <= ext.EXTERNAL_MAX_IMPORTANCE for e in ext_evs)
        assert any(e.url == "https://meti.example/cf2025" for e in evs)
        from deep_research_tool.utils.helpers import ResearchWarnings
        warns = " ".join(str(w) for w in ResearchWarnings.get_instance().get_all()) \
            if hasattr(ResearchWarnings.get_instance(), "get_all") else ""
        assert "壊れ" in json.dumps(res.get("warnings") or [], ensure_ascii=False) + warns

    def test_full_mode_cancel_and_no_stale_material(self, tmp_path):
        import threading
        from deep_research_tool.utils.cancellation import RunCancelled
        tool1, res1 = _run_v1(tmp_path)
        parent_path = res1["saved_artifacts"]["session"]
        f = tmp_path / "odr.md"
        f.write_text(ODR_MD, encoding="utf-8")
        # cancel before the work: nothing is compared, nothing fetched
        llm = ExtLLM()
        search = FakeSearchClient()
        tool = build_tool(tmp_path, llm, search, enable_verification=False, research_iterations=1)
        ev = threading.Event()
        ev.set()
        with pytest.raises(RunCancelled):
            tool.revise_report(parent_session_path=parent_path, feedback="", mode="add",
                               external_reports=[{"label": "ODR", "path": str(f)}], cancel_event=ev)
        assert not any("外部の調査資料（章と要旨・数値）" in p for p in llm.prompts)
        assert tool.cancel_event is ev
        # request_cancel() during the preparation reaches the same event
        tool = build_tool(tmp_path, ExtLLM(), FakeSearchClient(), enable_verification=False, research_iterations=1)
        orig = tool.search_client.get_page_content

        def cancel_on_fetch(url, *a, **kw):
            tool.request_cancel()
            return orig(url)
        tool.search_client.get_page_content = cancel_on_fetch
        with pytest.raises(RunCancelled):
            tool.revise_report(parent_session_path=parent_path, feedback="", mode="add",
                               external_reports=[{"label": "ODR", "path": str(f)}])
        # full mode: the comparison table reports outcomes on the new chapters
        tool2 = build_tool(tmp_path, ExtLLM(), FakeSearchClient(), enable_verification=False, research_iterations=1)
        tool2.config.research.use_enhanced_synthesis = False
        tool2.config.research.max_gap_fill_rounds = 0
        res = tool2.revise_report(parent_session_path=parent_path, feedback="", mode="full",
                                  external_reports=[{"label": "ODR", "path": str(f)}], external_fetch_urls=False)
        x = res["revision"]["external"]
        assert res["revision"]["mode"] == "full"
        outcomes = {t["topic"]: t["outcome"] for t in x["topics"]}
        assert outcomes["市場規模"].startswith("全面再調査") and outcomes["中国メーカーの台頭"].startswith("全面再調査")
        reqs = tool2.researcher.get_session().requirements
        assert "【外部の調査資料との照合結果】" in reqs and "【出典の書き方】" in reqs
        # revising that revision without new material does not inherit its material list
        tool3 = build_tool(tmp_path, ExtLLM(), FakeSearchClient(), enable_verification=False, research_iterations=1)
        tool3.config.research.use_enhanced_synthesis = False
        tool3.config.research.max_gap_fill_rounds = 0
        tool2b = build_tool(tmp_path, ExtLLM(), FakeSearchClient(), enable_verification=False, research_iterations=1)
        tool2b.config.research.use_enhanced_synthesis = False
        tool2b.config.research.max_gap_fill_rounds = 0
        res_add = tool2b.revise_report(parent_session_path=parent_path, feedback="", mode="add",
                                       external_reports=[{"label": "ODR", "path": str(f)}], external_fetch_urls=False)
        assert tool2b.researcher.get_session().run_config["external_reports"][0]["label"] == "ODR"
        res3 = tool3.revise_report(parent_session_path=res_add["saved_artifacts"]["session"],
                                   feedback=FEEDBACK, mode="add")
        assert "external_reports" not in tool3.researcher.get_session().run_config
        assert "external" not in res3["revision"]

    def test_options_off_and_errors(self, tmp_path):
        tool1, res1 = _run_v1(tmp_path)
        parent_path = res1["saved_artifacts"]["session"]
        f = tmp_path / "odr.md"
        f.write_text(ODR_MD, encoding="utf-8")
        search = FakeSearchClient()
        tool2 = build_tool(tmp_path, ExtLLM(), search, enable_verification=False, research_iterations=1)
        tool2.config.research.use_enhanced_synthesis = False
        tool2.config.research.max_gap_fill_rounds = 0
        res = tool2.revise_report(parent_session_path=parent_path, feedback=FEEDBACK, mode="revise",
                                  external_reports=[{"label": "ODR", "path": str(f)}],
                                  external_fetch_urls=False, external_body_as_evidence=False)
        x = res["revision"]["external"]
        assert x["cited_urls"] == [] and x["body_as_evidence"] is False and x["fetch_cited"] is False
        evs = tool2.researcher.get_evidence_locker().get_all_evidence()
        assert not [e for e in evs if e.access_method == "external"]
        assert not any(e.url == "https://meti.example/cf2025" for e in evs)
        assert "【出典の書き方】" not in tool2.researcher.get_session().requirements
        with pytest.raises(ValueError, match="外部の調査資料も読めませんでした"):
            tool2.revise_report(parent_session_path=parent_path, feedback=" ",
                                external_reports=[{"path": str(tmp_path / "none.md")}])
        with pytest.raises(ValueError, match="空です"):
            tool2.revise_report(parent_session_path=parent_path, feedback=" ")


# ----------------------------------------------------------------------
# Web API
# ----------------------------------------------------------------------

class TestWebApi(_RevWeb):
    def test_feedback_with_external_material(self, web, monkeypatch):
        self._offline(monkeypatch)
        import deep_research_tool.tests.test_revision as tr
        monkeypatch.setattr(tr, "RevLLM", ExtLLM)
        status, d = _post(web, "/api/research", {**self.BASE, "query": "炭素繊維の市場調査",
                                                 "requirements": "目次は2章構成でよい"})
        parent = self._wait(web, d["job_id"])
        assert parent["state"] == "completed", parent.get("error")
        out_dir = Path(web.output_dir)
        up = out_dir / "uploads" / "t1"
        up.mkdir(parents=True)
        (up / "odr.md").write_text(ODR_MD, encoding="utf-8")
        # precheck: chapters, size, cited URLs; a path outside the output dir is not read
        status, d = _post(web, "/api/precheck-external", {"external_reports": [
            {"label": "ODR", "path": str(up / "odr.md")}, {"path": "/etc/passwd"},
            {"label": "貼り付け", "text": "## a\n\nb https://p.example/"}]})
        assert status == 200
        m = d["materials"]
        assert m[0]["sections"] == 3 and m[0]["cited_urls"] == 3 and m[1]["error"]
        assert m[2]["cited_urls"] == 1
        # validation
        status, d = _post(web, "/api/feedback", {**self.BASE, "job_id": parent["job_id"], "feedback": "",
                                                 "external_reports": [{"path": "/etc/passwd"}]})
        assert status == 400 and d["field"] == "feedback_external"
        status, d = _post(web, "/api/feedback", {**self.BASE, "job_id": parent["job_id"], "feedback": "",
                                                 "external_reports": [{"path": str(up / "odr.md")}] * 6})
        assert status == 400 and "5 件まで" in d["error"]
        status, d = _post(web, "/api/feedback", {**self.BASE, "job_id": parent["job_id"], "feedback": "",
                                                 "external_reports": "x"})
        assert status == 400
        status, d = _post(web, "/api/feedback", {**self.BASE, "job_id": parent["job_id"], "feedback": "",
                                                 "external_reports": [{"path": 1}, {"path": ["a"]}]})
        assert status == 400 and d["field"] == "feedback_external"
        status, d = _post(web, "/api/precheck-external", {"external_reports": {"a": 1}})
        assert status == 400
        # a rejected request writes no pasted file (too many items / bad mode)
        status, d = _post(web, "/api/feedback", {**self.BASE, "job_id": parent["job_id"], "feedback": "",
                                                 "external_reports": [{"text": "x" * 1000}] * 2000})
        assert status == 400 and "5 件まで" in d["error"]
        status, d = _post(web, "/api/feedback", {**self.BASE, "job_id": parent["job_id"], "feedback": "",
                                                 "mode": "zzz", "external_reports": [{"text": "## a\n\nb"}]})
        assert status == 400 and d["field"] == "feedback_mode"
        assert not list(out_dir.glob("uploads/paste_*"))
        # material only (no feedback text): a file and pasted text
        status, d = _post(web, "/api/feedback", {**self.BASE, "job_id": parent["job_id"], "feedback": "",
                                                 "mode": "add", "external_fetch_urls": False,
                                                 "external_reports": [
                                                     {"label": "ODR", "path": str(up / "odr.md")},
                                                     {"label": "", "text": "## 補足\n\n追加の本文。"}]})
        assert status == 202, d
        child = self._wait(web, d["job_id"])
        assert child["state"] == "completed", child.get("error")
        x = child["result"]["revision"]["external"]
        assert [mm["label"] for mm in x["materials"]] == ["ODR", "貼り付けた資料2"]
        assert [mm["origin"] for mm in x["materials"]] == ["file", "paste"]
        assert x["materials"][0]["file"] == "odr.md" and x["materials"][1]["file"] is None
        assert x["fetch_cited"] is False
        spec = child["params_summary"]["feedback"]
        # pasted text was saved under uploads/: the ledger keeps the path only
        assert all("text" not in e for e in spec["external_reports"])
        pasted = Path(spec["external_reports"][1]["path"])
        assert pasted.is_file() and pasted.read_text(encoding="utf-8").startswith("## 補足")
        assert out_dir.resolve() in pasted.resolve().parents
