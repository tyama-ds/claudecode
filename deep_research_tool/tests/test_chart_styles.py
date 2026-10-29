"""
Chart styles (report.chart_style): "classic" keeps the v1.5 look and only
fixes defects; "enhanced" adds the recommended changes.

- setting: config / Web UI parameter validation, classic by default
- fixes for both styles: scale repair against the original text, year
  coercion, a point claimed through the locker's section map, legend
  names that name another series, y_max far above the data, the unit in
  the title, yield counters, unparsable LLM replies told apart from []
- rendering QA in both styles over awkward panels: negative values on the
  axis, no overlapping labels, nothing cut off at the image edge; the
  enhanced style draws at the printed width and turns crowded bars
  horizontal
- grounding: a value that is not in its quote / a quote not in the page
- facts: a judgment that contradicts the data is replaced; enhanced adds
  the computed fact and the data's source line
- enhanced data rules: research firms / PR wires are 二次, actual before
  forecast for one year
- PDF / HTML reports carry the figures under their chapters and still pass
  the frozen-body artifact check; the pipeline writes chart_yield.json
No paid API: fake clients only.
"""

import copy
import json
from pathlib import Path

import pytest

from deep_research_tool.chartkit import build_section_catalog, render_png, validate_panels
from deep_research_tool.chartkit.data_layer import (SourcePolicy, _choose, classify_source_basis,
                                                    coerce_year, repair_scale)
from deep_research_tool.chartkit.facts import (apply_facts, judgment_conflict, panel_facts,
                                               source_line)
from deep_research_tool.chartkit.grounding import (QUOTE_NOT_IN_SOURCE, VALUE_NOT_IN_QUOTE,
                                                   ground_panel, quote_in_source, value_in_quote)
from deep_research_tool.chartkit.spec import fix_title_unit, parse_panels
from deep_research_tool.evidence.numerical_extractor import NumericalDataPoint, NumericalDataStore

SAMPLE = Path(__file__).resolve().parent.parent / "chartkit" / "samples" / "zyn"
PUB = "公表"


def _dp(subject, metric, year, value, unit, url="https://example.com/a", **kw):
    kw.setdefault("raw_text", f"{value}{unit}")
    return NumericalDataPoint(value=value, unit=unit, metric_name=metric, subject=subject,
                              year=year, source_url=url, source_title=url.split("/")[2], **kw)


def _series(values, unit, basis=None, kind="comparison", name="", sources=None, labels_ref="L"):
    return {"unit": unit, "values": values, "basis": basis or [PUB] * len(values),
            "name": name, "kind": kind, "labels_ref": labels_ref, "source": "s",
            "sources": sources or [{"url": "https://www.meti.go.jp/x", "title": "経産省",
                                    "evidence_id": "e1", "raw_text": f"{v}{unit}"} for v in values]}


# --------------------------------------------------------------------------
# setting
# --------------------------------------------------------------------------

class TestSetting:

    def test_config_default_and_validation(self):
        from deep_research_tool.config import create_config
        assert create_config(provider="openai", openai_api_key="t").report.chart_style == "classic"
        cfg = create_config(provider="openai", openai_api_key="t", chart_style="Enhanced")
        assert cfg.report.chart_style == "enhanced"
        with pytest.raises(ValueError):
            create_config(provider="openai", openai_api_key="t", chart_style="fancy")

    def test_web_parameter(self):
        from deep_research_tool.webui.server import _CONFIG_PARAM_MAP, _chart_style_param
        assert _CONFIG_PARAM_MAP["chart_style"] == "chart_style"
        assert _chart_style_param(" enhanced ") == "enhanced"
        with pytest.raises(ValueError):
            _chart_style_param("zzz")

    def test_ui_selector_round_trip(self):
        page = (Path(__file__).resolve().parent.parent / "webui" / "static" / "index.html").read_text(
            encoding="utf-8")
        assert 'id="feat_chart_style"' in page and 'value="classic" selected' in page
        assert "s.chart_style = $('feat_chart_style').value" in page
        assert "p.chart_style = $('feat_chart_style').value" in page


# --------------------------------------------------------------------------
# fixes for both styles: data and spec
# --------------------------------------------------------------------------

class TestDataFixes:

    def test_repair_scale(self):
        # the number of '50億ドル' with a billion unit: 10x off → the text's scale
        assert repair_scale(50, "市場は50億ドル規模", "billion USD") == ("億ドル", True)
        # a correct conversion is left alone
        assert repair_scale(5, "市場は50億ドル規模", "billion USD") == ("billion USD", False)
        assert repair_scale(1.2, "売上高は1.2兆円", "億円") == ("兆円", True)
        # English scale word: the unit keeps its own currency spelling
        assert repair_scale(5, "about 5 billion dollars", "億ドル") == ("billion ドル", True)
        # compound scales (千万) and percentages are never touched
        assert repair_scale(3, "3千万円", "億円") == ("億円", False)
        assert repair_scale(40, "シェア40%", "%") == ("%", False)

    def test_coerce_year(self):
        assert coerce_year(2024) == 2024 and coerce_year("FY2023") == 2023
        assert coerce_year("2025年") == 2025 and coerce_year("n/a") is None
        assert coerce_year(True) is None and coerce_year(3000) is None

    def test_point_claimed_through_the_locker_section_map(self):
        pts = [_dp("市場", "規模", 2023, 10, "億円", section_id="2", evidence_id="e9"),
               _dp("市場", "規模", 2024, 12, "億円", section_id="2", evidence_id="e9")]
        assert build_section_catalog(pts, "1", []).is_empty()
        cat = build_section_catalog(pts, "1", ["e9"])
        assert not cat.is_empty()
        assert cat.stats["points"] == 2 and cat.stats["series"] == 1

    def test_yield_counters(self):
        pts = [_dp("A", "", 2024, 1, "億円", section_id="1"),           # no metric
               _dp("A", "売上", 2024, 5, "億円", section_id="1")]       # one year, one entity
        cat = build_section_catalog(pts, "1")
        assert cat.is_empty()
        assert cat.stats["no_metric"] == 1 and cat.stats["time_too_short"] == 1
        assert cat.stats["comparison_too_few"] == 1 and cat.stats["series"] == 0


class TestSpecFixes:

    def _data(self):
        return {"years": ["2023", "2024"],
                "ts1": dict(_series([10, 12], "億円", kind="time", name="A社 売上", labels_ref="years"),
                            entity="a社"),
                "ts2": dict(_series([30, 25], "億円", kind="time", name="B社 売上", labels_ref="years"),
                            entity="b社")}

    def _panel(self, **kw):
        p = {"id": "p1", "form": "bar", "title": "売上（億円）", "labels_ref": "years",
             "series": [{"name": "B社", "values_ref": "ts1", "role": "emphasis"},
                        {"name": "A社", "values_ref": "ts2", "role": "context"}],
             "caption": "c", "judgment": "j"}
        p.update(kw)
        return p

    def test_swapped_legend_names_take_the_data_names(self):
        valid, rej = validate_panels([self._panel()], self._data())
        assert not rej
        names = [s["name"] for s in valid[0]["series"]]
        assert names == ["A社 売上", "B社 売上"] and len(valid[0]["name_fixes"]) == 2

    def test_y_max_far_above_the_data_is_dropped(self):
        valid, _ = validate_panels([self._panel(y_max=5000)], self._data())
        assert "y_max" not in valid[0]
        valid, _ = validate_panels([self._panel(y_max=40)], self._data())
        assert valid[0]["y_max"] == 40

    def test_title_unit(self):
        assert fix_title_unit("市場規模（億円）", "百万円") == "市場規模（百万円）"
        assert fix_title_unit("推計：調査会社別（10億ドル）", "億ドル") == "推計：調査会社別（億ドル）"
        assert fix_title_unit("推計（調査会社別）", "億ドル") == "推計（調査会社別）"
        assert fix_title_unit("推計", "億ドル") == "推計"
        assert fix_title_unit("推計", "億ドル", append=True) == "推計（億ドル）"

    def test_unparsable_reply_told_apart_from_an_empty_list(self):
        st = {}
        assert parse_panels("作れません", st) == [] and st["ok"] is False
        assert parse_panels("[]", st) == [] and st["ok"] is True

    def test_over_limit_is_counted(self):
        p2 = self._panel(id="p2", form="line_xy")
        stats = {}
        valid, rej = validate_panels([self._panel(), p2], self._data(), max_panels=1, stats=stats)
        assert len(valid) == 1 and stats["over_limit"] == 1

    def test_line_xy_sources_stay_aligned_with_points(self):
        data = {"years": ["2022", "2023", "2024"],
                "ts1": _series([5, None, 7], "億円", basis=[PUB, None, PUB], kind="time",
                               labels_ref="years",
                               sources=[{"raw_text": "5億円"}, None, {"raw_text": "7億円"}])}
        p = {"id": "x", "form": "line_xy", "title": "t", "labels_ref": "years",
             "series": [{"name": "s", "values_ref": "ts1", "role": "emphasis"}],
             "caption": "c", "judgment": "j"}
        valid, _ = validate_panels([p], data)
        s = valid[0]["series"][0]
        assert [v["y"] for v in s["values"]] == [5.0, 7.0]
        assert [x["raw_text"] for x in s["sources"]] == ["5億円", "7億円"]


# --------------------------------------------------------------------------
# rendering QA, both styles
# --------------------------------------------------------------------------

def _qa_panels():
    long = [f"製品カテゴリ{i}（国内向け）" for i in range(12)]
    data = {
        "cats": ["A社", "B社", "C社", "D社", "E社"],
        "long": long,
        "years": ["2019", "2020", "2021", "2022", "2023"],
        "neg": _series([12, -8, 5, -15, 3], "%", labels_ref="cats"),
        "many": _series([1200 + 137 * i for i in range(12)], "億円", labels_ref="long"),
        "growth": _series([5, -12, 8, 3, -2], "%", kind="time", labels_ref="years"),
        "close1": _series([9.0, 9.4, 9.5, 9.6, 10], "億円", kind="time", labels_ref="years"),
        "close2": _series([9.1, 9.3, 9.6, 9.5, 10.2], "億円", kind="time", labels_ref="years"),
        "tiny": _series([0.04, 0.031, 0.05, 0.02, 0.06], "億円", labels_ref="cats"),
        "huge": _series([1_250_000, 980_000, 2_400_000, 1_700_000, 640_000], "トン", labels_ref="cats"),
        "segA": _series([56, 40, 30, 1, 45], "%", labels_ref="cats"),
        "segB": _series([3, 35, 50, 60, 10], "%", labels_ref="cats"),
        "segC": _series([41, 25, 20, 39, 45], "%", labels_ref="cats"),
    }

    def panel(pid, form, ref, keys, **kw):
        roles = ["emphasis"] + ["context"] * 7 if not form.startswith("stacked") else ["part"] * 4
        p = {"id": pid, "form": form, "title": f"{pid}（単位）", "labels_ref": ref,
             "series": [{"name": k, "values_ref": k, "role": r} for k, r in zip(keys, roles)],
             "caption": "出所: テスト（公表）。定義の違いに注意。" * 2,
             "judgment": "比較して見る。"}
        p.update(kw)
        return p
    panels = [
        panel("neg_hbar", "hbar", "cats", ["neg"]),
        panel("neg_bar", "bar", "cats", ["neg"]),
        panel("many_bar", "bar", "long", ["many"]),
        panel("growth_line", "line_xy", "years", ["growth"]),
        panel("close_lines", "line_xy", "years", ["close1", "close2"]),
        panel("tiny_bar", "bar", "cats", ["tiny"]),
        panel("huge_hbar", "hbar", "cats", ["huge"]),
        panel("segments", "stacked_hbar", "cats", ["segA", "segB", "segC"]),
        panel("share_ymax", "hbar", "cats", ["segB"], y_max=100),
    ]
    valid, rejected = validate_panels(panels, data)
    assert not rejected, rejected
    return valid


@pytest.mark.parametrize("style", ["classic", "enhanced"])
def test_render_qa(tmp_path, style):
    pytest.importorskip("matplotlib")
    for panel in _qa_panels():
        qa = {}
        render_png(copy.deepcopy(panel), tmp_path / f"{panel['id']}.png", style=style, qa=qa)
        pid = panel["id"]
        assert qa["label_overlaps"] == 0, (pid, qa)
        assert qa["tick_overlaps"] == 0, (pid, qa)
        assert qa["spill"] == [], (pid, qa["spill"])
        lo, hi = qa["value_lim"]
        vals = [v["y"] if isinstance(v, dict) else v for s in panel["series"] for v in s["values"]
                if v is not None]
        if not panel["form"].startswith("stacked"):
            assert lo <= min(vals) and max(vals) <= hi, (pid, qa["value_lim"])
        if style == "enhanced":
            assert qa["size_px"][0] == 1100           # 5.5 in at 200 dpi: printed 1:1
        else:
            assert qa["size_px"][0] == 1500           # the v1.5 size


def test_enhanced_turns_crowded_bars_horizontal_classic_keeps_the_form(tmp_path):
    pytest.importorskip("matplotlib")
    many = next(p for p in _qa_panels() if p["id"] == "many_bar")
    c, e = copy.deepcopy(many), copy.deepcopy(many)
    render_png(c, tmp_path / "c.png", style="classic")
    render_png(e, tmp_path / "e.png", style="enhanced")
    assert "form_rendered" not in c and e["form_rendered"] == "hbar"


def test_value_labels_keep_small_numbers_and_separators():
    from deep_research_tool.chartkit.render_png import _fmt, _wrap
    assert _fmt(0.04, "億円") == "0.04" and _fmt(2707, "億円") == "2,707"
    assert _fmt(-8, "%", signed=True) == "-8%"
    wrapped = _wrap("2024年は6億缶、1.5倍に増えた。" * 3, 12)
    assert "6億缶" in wrapped.replace("\n", "") and all(not ln.startswith("、") for ln in wrapped.split("\n"))
    assert "1.5" in wrapped and "1.\n5" not in wrapped


# --------------------------------------------------------------------------
# grounding
# --------------------------------------------------------------------------

class TestGrounding:

    def test_value_and_quote(self):
        assert value_in_quote(3.1, "10億ドル", "約31億ドル")
        assert value_in_quote(56, "%", "数量シェア56%")
        assert value_in_quote(-8, "%", "前年比8%減")
        assert not value_in_quote(5, "億円", "売上高は50億円")
        assert quote_in_source("売上高は1,234億円", "当社の売上高は 1234 億円でした")
        assert quote_in_source("2024年 売上 1,234億円", "…2024年度の売上は1,234億円…")
        assert not quote_in_source("売上高は999億円", "当社の売上高は1234億円でした")

    def test_ground_panel_flags_both_problems(self):
        srcs = [{"evidence_id": "e1", "url": "u1", "raw_text": "A社は120億円"},
                {"evidence_id": "e2", "url": "u2", "raw_text": "B社は80億円"},
                {"evidence_id": "e3", "url": "u3", "raw_text": "C社は30億円"}]
        panel = {"labels": ["A社", "B社", "C社"], "unit": "億円",
                 "series": [{"name": "売上", "unit": "億円", "values": [120, 8, 30],
                             "basis": [PUB, "二次", "導出"], "sources": srcs}]}
        pages = {"e1": "決算: A社は120億円", "e2": "B社は80億円"}
        g = ground_panel(panel, lambda eid, url: pages.get(eid))
        assert g["checked"] == 2                    # 導出 is computed, not quoted
        assert [(i["label"], i["problem"]) for i in g["issues"]] == [("B社", VALUE_NOT_IN_QUOTE)]
        pages["e1"] = "別の記事"
        g = ground_panel(panel, lambda eid, url: pages.get(eid))
        assert ("A社", QUOTE_NOT_IN_SOURCE) in [(i["label"], i["problem"]) for i in g["issues"]]


# --------------------------------------------------------------------------
# facts
# --------------------------------------------------------------------------

def _line(values, judgment, forecast_last=False):
    srcs = [{"title": "業界団体", "url": "https://x.or.jp/a", "forecast": False} for _ in values]
    if forecast_last:
        srcs[-1] = dict(srcs[-1], forecast=True)
    return {"form": "line_xy", "unit": "%", "labels": [], "title": "成長率", "caption": "業界団体",
            "judgment": judgment,
            "series": [{"name": "成長率", "unit": "%", "basis": [PUB] * len(values), "sources": srcs,
                        "values": [{"x": float(2019 + i), "y": v} for i, v in enumerate(values)]}]}


class TestFacts:

    def test_only_plain_max_min_claims_are_checked(self):
        f = panel_facts(_line([5, -12, 8, 3, -2], ""))
        assert f["kind"] == "time" and f["direction"] == "down"
        # direction words cannot be read reliably (subject, cause, rate…): never checked
        for t in ("一貫して増加している。", "2020年に大きく落ち込んだ。", "2020年に増加した。",
                  "2020年に落ち込んだが2021年に回復。", "一貫した増加は見られない。"):
            assert judgment_conflict(t, f) is None, t
        bar = {"form": "bar", "unit": "億円", "labels": ["A社", "B社", "C社"],
               "series": [{"values": [300, 120, 50]}]}
        fb = panel_facts(bar)
        assert judgment_conflict("B社が最大。", fb) and judgment_conflict("A社が最大。", fb) is None

    def test_the_judgment_is_never_rewritten(self):
        ok = _line([5, -12, 8, 3, -2], "2020年に大きく落ち込んだ。")
        assert apply_facts(ok, "classic") == {} and ok["judgment"] == "2020年に大きく落ち込んだ。"
        assert ok["caption"] == "業界団体" and ok["title"] == "成長率"
        bad = _line([5, -12, 8, 3, -2], "2019年が過去最高。")      # 2021 is higher
        rec = apply_facts(bad, "classic")
        assert rec["conflict"]["judgment"] == "2019年が過去最高。" and bad["judgment"] == "2019年が過去最高。"
        bad2 = _line([5, -12, 8, 3, -2], "2019年が過去最高。")
        rec2 = apply_facts(bad2, "enhanced")
        assert rec2["conflict"] and "judgment_replaced" not in rec2
        # enhanced: the sentence stays, the computed facts follow it
        assert bad2["judgment"].startswith("2019年が過去最高。データ: 2019→2023で5%→-2%（減少")

    def test_enhanced_adds_fact_source_line_and_unit(self):
        p = _line([5, -12, 8, 3, -2], "2020年に大きく落ち込んだ。", forecast_last=True)
        rec = apply_facts(p, "enhanced")
        assert p["judgment"] == "2020年に大きく落ち込んだ。データ: 2019→2023（予測）で5%→-2%（減少、-7pt）。"
        assert "出所（データ）: 業界団体（公表5）" in p["caption"] and "予測値: 2023" in p["caption"]
        assert p["title"] == "成長率（%）" and rec["source_line"]

    def test_source_line_counts_and_truncates(self):
        srcs = [{"title": f"資料{i}", "url": f"https://s{i}.example/x"} for i in range(5)]
        panel = {"labels": list("abcde"),
                 "series": [{"values": [1, 2, 3, 4, 5], "basis": [PUB, PUB, "二次", "換算", "二次"],
                             "sources": srcs}]}
        assert source_line(panel) == "出所（データ）: 資料0、資料1、資料2 ほか2件（公表2・換算1・二次2）。"


# --------------------------------------------------------------------------
# enhanced data rules
# --------------------------------------------------------------------------

class TestEnhancedSourceRules:

    def test_research_firms_and_pr_wires_are_secondary(self):
        classic = SourcePolicy.for_style("classic")
        enhanced = SourcePolicy.for_style("enhanced")
        pr = "https://prtimes.jp/main/html/rd/p/000.html"
        firm = "https://www.grandviewresearch.com/press-release/global-market"
        assert classify_source_basis(pr, classic.primary_domains, classic.secondary_domains) == "公表"
        assert classify_source_basis(pr, enhanced.primary_domains, enhanced.secondary_domains) == "二次"
        assert classify_source_basis(firm, classic.primary_domains) == "公表"   # the /press-release path
        assert classify_source_basis(firm, enhanced.primary_domains, enhanced.secondary_domains) == "二次"
        assert classify_source_basis("https://www.meti.go.jp/x", enhanced.primary_domains,
                                     enhanced.secondary_domains) == "公表"

    def test_actual_before_forecast_for_one_year(self):
        gov = "https://www.stat.go.jp/x"
        forecast = _dp("市場", "規模", 2024, 120, "億円", gov, is_forecast=True)
        actual = _dp("市場", "規模", 2024, 110, "億円", "https://news.example.com/a")
        assert _choose([forecast, actual], SourcePolicy.for_style("classic")) is forecast
        assert _choose([forecast, actual], SourcePolicy.for_style("enhanced")) is actual


# --------------------------------------------------------------------------
# PDF / HTML reports
# --------------------------------------------------------------------------

class TestReportFigures:

    def _setup(self, tmp_path):
        from deep_research_tool.report.figure_table_generator import (
            Figure, FigureTableCollection, FigureType, TableData)
        from deep_research_tool.report.semantic_manifest import figure_semantics_markdown
        from deep_research_tool.tests.test_finalization_integration import (
            make_locker, make_plan, make_session)
        data = json.loads((SAMPLE / "data.json").read_text(encoding="utf-8"))
        panels = json.loads((SAMPLE / "panels.json").read_text(encoding="utf-8"))
        valid, _ = validate_panels(panels, data)
        png = tmp_path / "chart.png"
        render_png(valid[0], png, style="enhanced")
        coll = FigureTableCollection(
            charts=[Figure(figure_id="c1", figure_type=FigureType.CHART, title=valid[0]["title"],
                           caption="出所 2 件　見るべき点: 回復基調。", section_id="1", image_path=png),
                    Figure(figure_id="c2", figure_type=FigureType.CHART, title="章外の図",
                           caption="どの章にも属さない図。", section_id="9", image_path=png)],
            tables=[TableData(table_id="t1", title="主要メーカー", caption="出所: 各社IR",
                              headers=["社名", "シェア"], rows=[["A社", "40%"], ["B社", "25%"]],
                              section_id="2")])
        plan = make_plan([("1", "市場規模の推移", "d"), ("2", "競合", "d")])
        ev = [("https://ex.com/a", "x " * 30)]
        session = make_session(plan, {"1": {"title": "市場規模の推移", "content": "本文 [1]", "evidence": ev},
                                      "2": {"title": "競合", "content": "本文2 [1]", "evidence": ev}})
        appendix = figure_semantics_markdown(coll, language="ja", section_id="付録D")
        session.section_contents["付録D"] = {"title": "付録D", "content": appendix,
                                            "sources": [], "extracted_content": []}
        return session, make_locker(tmp_path, ev), coll, appendix

    def test_pdf_places_figures_under_their_chapters(self, tmp_path):
        pypdf = pytest.importorskip("pypdf")
        pytest.importorskip("reportlab")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.generator import ReportGenerator
        from deep_research_tool.report.semantic_manifest import build_semantic_manifest, check_artifact
        session, locker, coll, appendix = self._setup(tmp_path)
        gen = ReportGenerator(output_dir=tmp_path / "reports", language="ja")
        gen.figure_collection = coll
        path = gen.generate_report(session=session, evidence_locker=locker,
                                   format=ReportFormat.PDF, verification_result=None)
        assert gen.figures_inserted == 3
        reader = pypdf.PdfReader(str(path))
        text = "\n".join(p.extract_text() for p in reader.pages)
        assert sum(len(p.images) for p in reader.pages) == 2
        assert text.index("1. 市場規模の推移") < text.index("見るべき点: 回復基調") < text.index("2. 競合")
        assert text.index("2. 競合") < text.index("主要メーカー")
        assert "どの章にも属さない図" in text                     # never silently dropped
        manifest = build_semantic_manifest({"1": "本文 [1]", "2": "本文2 [1]", "付録D": appendix},
                                           figure_collection=coll)
        assert check_artifact(manifest, path, locker)["status"] == "pass"

    def test_html_embeds_figures(self, tmp_path):
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.generator import ReportGenerator
        session, locker, coll, _ = self._setup(tmp_path)
        gen = ReportGenerator(output_dir=tmp_path / "reports", language="ja")
        gen.figure_collection = coll
        path = gen.generate_report(session=session, evidence_locker=locker,
                                   format=ReportFormat.HTML, verification_result=None)
        page = path.read_text(encoding="utf-8")
        assert page.count('class="report-figure"><img src="data:image/png;base64,') == 2
        # (the table of contents lists the chapters first: anchor on the headings)
        assert page.index("<h2>1. 市場規模の推移</h2>") < page.index("見るべき点: 回復基調") \
            < page.index("<h2>2. 競合</h2>")
        assert "<caption>主要メーカー — 出所: 各社IR</caption>" in page

    def test_html_report_passes_the_artifact_check(self, tmp_path):
        # the page's own CSS used to count as unverified additions, so every
        # verified V1 HTML report failed the check
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.generator import ReportGenerator
        from deep_research_tool.report.semantic_manifest import build_semantic_manifest, check_artifact
        session, locker, coll, appendix = self._setup(tmp_path)
        gen = ReportGenerator(output_dir=tmp_path / "reports", language="ja")
        gen.figure_collection = coll
        path = gen.generate_report(session=session, evidence_locker=locker,
                                   format=ReportFormat.HTML, verification_result=None)
        manifest = build_semantic_manifest({"1": "本文 [1]", "2": "本文2 [1]", "付録D": appendix},
                                           figure_collection=coll)
        res = check_artifact(manifest, path, locker)
        assert res["status"] == "pass", res

    def test_pdf_bullets_of_a_long_limitation_are_not_additions(self, tmp_path):
        # the limitations land in the LAST chapter — with figures, the
        # appendix — whose "- " items the PDF draws as "• "
        pytest.importorskip("reportlab")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.generator import ReportGenerator
        from deep_research_tool.report.semantic_manifest import (build_semantic_manifest, check_artifact,
                                                                 normalize_semantic_text)
        assert normalize_semantic_text("• 未解決: 範囲") == normalize_semantic_text("- 未解決: 範囲")
        session, locker, coll, appendix = self._setup(tmp_path)
        appendix += ("\n### 調査上の限界\n\n- 未解決: 未検証の本文範囲が残っています（抽出チャンク 1/4 失敗、"
                     "リトライ2回後）。この範囲の数値は出典で確認できていないため、判断には使わないでください。\n")
        session.section_contents["付録D"]["content"] = appendix
        gen = ReportGenerator(output_dir=tmp_path / "reports", language="ja")
        gen.figure_collection = coll
        path = gen.generate_report(session=session, evidence_locker=locker,
                                   format=ReportFormat.PDF, verification_result=None)
        manifest = build_semantic_manifest({"1": "本文 [1]", "2": "本文2 [1]", "付録D": appendix},
                                           figure_collection=coll)
        res = check_artifact(manifest, path, locker)
        assert res["status"] == "pass", res

    def test_without_a_collection_nothing_changes(self, tmp_path):
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.generator import ReportGenerator
        session, locker, _coll, _ = self._setup(tmp_path)
        gen = ReportGenerator(output_dir=tmp_path / "reports", language="ja")
        path = gen.generate_report(session=session, evidence_locker=locker,
                                   format=ReportFormat.HTML, verification_result=None)
        assert gen.figures_inserted == 0
        assert 'class="report-figure"' not in path.read_text(encoding="utf-8")


class TestInsertionPaths:

    def _tool(self, tmp_path):
        from deep_research_tool.config import create_config
        from deep_research_tool.main import DeepResearchTool
        tool = DeepResearchTool.__new__(DeepResearchTool)
        tool.config = create_config(provider="openai", openai_api_key="t", output_dir=str(tmp_path))
        return tool

    def test_docx_failure_warns_instead_of_crashing(self, tmp_path):
        from types import SimpleNamespace
        from deep_research_tool.report.figure_table_generator import FigureTableCollection
        from deep_research_tool.utils.helpers import ResearchWarnings
        gen = SimpleNamespace(output_dir=tmp_path, add_figures_to_docx=lambda **kw: None)
        report = tmp_path / "r.docx"
        report.write_bytes(b"x")
        collector = ResearchWarnings.bind()
        try:
            out = self._tool(tmp_path)._insert_figures_into_report(
                report_path=report, collection=FigureTableCollection(), generator=gen)
            assert out is None
            assert any("DOCXへの挿入に失敗" in w["message"] for w in collector.to_dict_list())
        finally:
            ResearchWarnings.unbind()

    def test_a_finished_pdf_is_not_replaced_by_a_figures_only_file(self, tmp_path):
        from types import SimpleNamespace
        from deep_research_tool.report.figure_table_generator import FigureTableCollection
        from deep_research_tool.utils.helpers import ResearchWarnings
        called = []
        gen = SimpleNamespace(output_dir=tmp_path,
                              add_figures_to_pdf=lambda **kw: called.append(kw) or tmp_path / "x.pdf")
        report = tmp_path / "r.pdf"
        report.write_bytes(b"%PDF")
        collector = ResearchWarnings.bind()
        try:
            out = self._tool(tmp_path)._insert_figures_into_report(
                report_path=report, collection=FigureTableCollection(), generator=gen)
            assert out is None and not called
            assert any("PDF には図表を挿入できません" in w["message"] for w in collector.to_dict_list())
        finally:
            ResearchWarnings.unbind()


# --------------------------------------------------------------------------
# pipeline: style, grounding, facts, yield
# --------------------------------------------------------------------------

class _Resp:
    def __init__(self, content):
        self.content = content


class _LLM:
    """Writes one panel per chapter whose judgment contradicts the data."""

    def __init__(self):
        self.calls = 0

    def generate(self, prompt, system_prompt=None, **kw):
        self.calls += 1
        cat = json.loads(prompt.split("値は含めない）:\n", 1)[1])
        it = next(c for c in cat if c["kind"] == "時系列")
        return _Resp(json.dumps([{
            "id": "p1", "form": "bar", "title": f"{it['name']}（{it['unit']}）",
            "labels_ref": it["labels_ref"],
            "series": [{"name": it["name"], "values_ref": it["values_ref"], "role": "emphasis"}],
            "caption": "出所: 統計", "judgment": "2021年が過去最高。"}], ensure_ascii=False))


@pytest.mark.parametrize("style", ["classic", "enhanced"])
def test_pipeline_style_grounding_facts_and_yield(tmp_path, style):
    pytest.importorskip("matplotlib")
    from types import SimpleNamespace
    from deep_research_tool.config import create_config
    from deep_research_tool.evidence.locker import EvidenceLocker
    from deep_research_tool.main import DeepResearchTool
    from deep_research_tool.utils.helpers import ResearchWarnings
    cfg = create_config(provider="openai", openai_api_key="t", output_dir=str(tmp_path),
                        auto_figures=True, chart_style=style)
    tool = DeepResearchTool.__new__(DeepResearchTool)
    tool.config, tool.llm_client, tool.stage_llm_clients = cfg, _LLM(), {}
    store = NumericalDataStore(research_topic="炭素繊維")
    gov = "https://www.stat.go.jp/x"
    for y, v in ((2021, 40), (2022, 50), (2023, 55)):
        store.add(_dp("炭素繊維市場", "市場規模", y, v, "億円", gov, section_id="1"))
    # a value whose quote says something else (grounding)
    store.add(_dp("炭素繊維市場", "市場規模", 2024, 60, "億円", gov, section_id="1",
                  raw_text="2024年は6.5億円"))
    locker = EvidenceLocker(research_id="r", output_dir=tmp_path / "ev")
    session = SimpleNamespace(query="炭素繊維", research_plan=None,
                              section_contents={"1": {"title": "市場規模", "content": "本文", "sources": []},
                                                "2": {"title": "需要", "content": "本文", "sources": []}})
    collector = ResearchWarnings.bind()
    try:
        figs = tool._spec_charts(session, locker, store, tmp_path / "figures")
        msgs = [w["message"] for w in collector.to_dict_list()]
    finally:
        ResearchWarnings.unbind()
    assert len(figs) == 1
    sem = figs[0].chart_data
    # the judgment is never rewritten; the contradiction is reported in both styles
    assert sem["judgment"].startswith("2021年が過去最高。")
    assert any("食い違う可能性" in m for m in msgs)
    if style == "enhanced":
        assert "データ: 2021→2024で40億円→60億円（増加" in sem["judgment"]
    else:
        assert sem["judgment"] == "2021年が過去最高。"
    assert any("出典本文で確認できません" in m for m in msgs)
    y = json.loads((tmp_path / "figures" / "chart_yield.json").read_text(encoding="utf-8"))
    assert y["style"] == style and y["charts"] == 1
    ch = y["chapters"]["1"]
    assert ch["catalog"]["points"] == 4 and ch["valid"] == 1 and ch["rendered"] == 1
    assert ch["panels"][0]["grounding"]["issues"][0]["label"] == "2024"
    assert y["chapters"]["2"]["catalog"]["points"] == 0
    if style == "enhanced":
        assert "出所（データ）: www.stat.go.jp" in sem["caption"] or "出所（データ）: stat.go.jp" in sem["caption"]
        assert "要確認" in sem["caption"]
    else:
        assert sem["caption"] == "出所: 統計"


# --------------------------------------------------------------------------
# regressions from the independent review
# --------------------------------------------------------------------------

class TestReviewRegressions:

    def test_title_parenthesis_that_is_not_a_unit_is_kept_over_percent(self):
        assert fix_title_unit("シェア（2024年）", "%") == "シェア（2024年）"
        assert fix_title_unit("営業利益率（主要5社）", "%") == "営業利益率（主要5社）"
        assert fix_title_unit("成長率（前年比）", "%") == "成長率（前年比）"
        assert fix_title_unit("売上（億円）", "%") == "売上（%）"

    def _line(self, vals, unit="億円", start=2017):
        return panel_facts({"form": "line_xy", "unit": unit, "labels": [], "series": [
            {"values": [{"x": float(start + i), "y": v} for i, v in enumerate(vals)]}]})

    def test_true_judgments_are_never_flagged(self):
        assert judgment_conflict("営業赤字は縮小が続いている", self._line([-120, -90, -60, -20])) is None
        f = self._line([100, 90, 80, 85, 95, 110], start=2016)
        assert judgment_conflict("2018年以降、増加が続いている", f) is None
        assert judgment_conflict("2018年比で増加", f) is None
        assert judgment_conflict("2022年は前年比の落ち込みが過去最大",
                                 self._line([100, 120, 90, 130], start=2020)) is None
        assert judgment_conflict("2021年にピークを迎え、その後減少",
                                 self._line([100, 140, 120, 110], start=2020)) is None
        assert judgment_conflict("2020年に落ち込んだが、その後回復",
                                 self._line([5, -12, 8, 3, -2], unit="%", start=2019)) is None
        assert judgment_conflict("2021年は増加、2022年も増加",
                                 self._line([100, 110, 120, 115], start=2020)) is None
        assert judgment_conflict("2023年は減少", self._line([100, 100, 101, 100.5], start=2020)) is None

    def test_plain_false_judgments_are_flagged(self):
        rising = self._line([100, 110, 120, 130, 140], start=2019)
        assert judgment_conflict("2020年が過去最高", rising)
        assert judgment_conflict("2020年は過去最高です。", rising)
        bar = panel_facts({"form": "bar", "unit": "億円", "labels": ["A社", "B社", "C社"],
                           "series": [{"values": [300, 120, 50]}]})
        assert judgment_conflict("B社が首位", bar) and judgment_conflict("A社が最も低い", bar)

    def test_more_true_judgments_are_not_flagged(self):
        bar = panel_facts({"form": "bar", "unit": "億円", "labels": ["A社", "B社", "C社"],
                           "series": [{"values": [300, 120, 50]}]})
        for t in ("B社は最大手A社を追う", "C社は最大の成長余地を持つ", "B社はトップとの差を縮めている",
                  "B社はトップに次ぐ2位", "A社が最大。", "C社が最も低い"):
            assert judgment_conflict(t, bar) is None, t
        line = self._line([100, 140, 120, 110], start=2018)
        for t in ("2019年をピークに減少", "A社のシェアは2019年に首位に浮上した", "2020年代は拡大が見込まれる",
                  "原料高により一貫して上昇", "2018年以降一貫して増加", "足元は減少している"):
            assert judgment_conflict(t, line) is None, t
        # 過去最高 of the latest actual year is not contradicted by a forecast year
        fc = panel_facts({"form": "line_xy", "unit": "億円", "labels": [], "series": [{
            "values": [{"x": 2022.0, "y": 100}, {"x": 2023.0, "y": 120}, {"x": 2024.0, "y": 150}],
            "sources": [{}, {}, {"forecast": True}]}]})
        assert judgment_conflict("2023年は過去最高", fc) is None

    def test_a_number_unit_reads_naturally_in_the_fact(self):
        from deep_research_tool.chartkit.facts import fact_sentence
        f = self._line([5, 6, 6.5, 7.1], unit="10億ドル", start=2021)
        assert "50億ドル→71億ドル" in fact_sentence(f) and "510" not in fact_sentence(f)

    def test_title_unit_is_not_added_twice(self):
        assert fix_title_unit("売上（億円）の推移", "億円", append=True) == "売上（億円）の推移"
        assert fix_title_unit("売上（百万円）の推移", "億円", append=True) == "売上（億円）の推移"

    def test_grounding_respects_the_scale(self):
        assert not value_in_quote(50, "10億ドル", "市場規模は50億ドル")      # 10x
        assert not value_in_quote(500, "億円", "500万円")                   # 10,000x
        assert value_in_quote(1234, "百万円", "売上高 1,234")               # scale in a table header
        assert not quote_in_source("売上高は50億円", "2050年に向けて150社")
        assert not quote_in_source("3.5%", "13.5%")

    def test_grounding_uses_the_extracted_value_for_a_bare_quote(self):
        panel = {"labels": ["2024"], "unit": "10億ドル", "series": [{
            "unit": "10億ドル", "values": [5.0], "basis": [PUB],
            "sources": [{"raw_text": "5,000", "value": 5000, "unit": "百万ドル"}]}]}
        assert ground_panel(panel, lambda e, u: None)["issues"] == []

    def test_repair_scale_scale_only_unit_and_negative_values(self):
        assert repair_scale(50, "50億ドル", "billion") == ("億", True)
        assert repair_scale(-50, "-50億ドル", "billion USD") == ("億ドル", True)

    def test_host_with_login_and_port(self):
        assert classify_source_basis("https://user:pw@www.meti.go.jp/x") == "公表"
        assert classify_source_basis("https://www.meti.go.jp:8443/x") == "公表"

    def test_negative_parts_cannot_be_stacked(self):
        data = {"cats": ["A", "B"], "a": _series([10, -5], "億円", labels_ref="cats"),
                "b": _series([3, 4], "億円", labels_ref="cats")}
        p = {"id": "s", "form": "stacked_bar", "title": "t", "labels_ref": "cats",
             "series": [{"name": "a", "values_ref": "a", "role": "part"},
                        {"name": "b", "values_ref": "b", "role": "part"}],
             "caption": "c", "judgment": "j"}
        valid, rej = validate_panels([p], data)
        assert not valid and "負の値" in rej[0][1]

    def test_wrap_cuts_a_token_longer_than_a_line(self):
        from deep_research_tool.chartkit.render_png import _wrap
        url = "https://www.meti.go.jp/" + "x" * 120
        assert all(len(line) <= 46 for line in _wrap("出所: " + url, 46).split("\n"))

    def test_classic_keeps_the_category_label_size(self):
        from deep_research_tool.chartkit.render_png import _LAYOUT
        assert _LAYOUT["classic"]["cat"] == 10 and _LAYOUT["classic"]["width"] == 10.0


def _tick_texts(panel, style, tmp_path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    seen = {}
    orig = plt.Figure.savefig

    def spy(self, *a, **k):
        ax = self.axes[0]
        lo, hi = ax.get_ylim()
        seen["ticks"] = [(loc, t.get_text()) for loc, t in zip(ax.yaxis.get_ticklocs(), ax.get_yticklabels())
                         if lo <= loc <= hi]
        seen["texts"] = [t.get_text() for t in ax.texts if t.get_visible()]
        seen["colors"] = {t.get_text(): t.get_position() for t in ax.texts}
        return orig(self, *a, **k)
    plt.Figure.savefig = spy
    try:
        render_png(copy.deepcopy(panel), tmp_path / f"t_{style}.png", style=style)
    finally:
        plt.Figure.savefig = orig
    return seen


def test_enhanced_ticks_print_the_tick_value(tmp_path):
    pytest.importorskip("matplotlib")
    data = {"cats": ["A", "B", "C"], "v": _series([0.5, 1.6, 1.9], "億円", labels_ref="cats")}
    p = {"id": "x", "form": "bar", "title": "t", "labels_ref": "cats",
         "series": [{"name": "v", "values_ref": "v", "role": "emphasis"}], "caption": "c", "judgment": "j"}
    valid, _ = validate_panels([p], data)
    ticks = _tick_texts(valid[0], "enhanced", tmp_path)["ticks"]
    for loc, text in ticks:
        assert abs(float(text.replace(",", "")) - loc) < 1e-9, (loc, text)


@pytest.mark.parametrize("style", ["classic", "enhanced"])
def test_big_values_and_long_names_stay_inside_the_image(tmp_path, style):
    pytest.importorskip("matplotlib")
    data = {"cats": ["A", "B", "C"], "long": ["非常に長いカテゴリ名の会社ホールディングス株式会社グループ全体の連結"] * 1
            + ["B社", "C社"],
            "big": _series([12_345_678, 23_456_789, 34_567_890], "億円", labels_ref="cats"),
            "lv": _series([3, 2, 1], "億円", labels_ref="long")}
    panels = [{"id": "big", "form": "bar", "title": "t", "labels_ref": "cats",
               "series": [{"name": "big", "values_ref": "big", "role": "emphasis"}],
               "caption": "c", "judgment": "j"},
              {"id": "long", "form": "hbar", "title": "t", "labels_ref": "long",
               "series": [{"name": "lv", "values_ref": "lv", "role": "emphasis"}],
               "caption": "c", "judgment": "j"}]
    valid, rej = validate_panels(panels, data)
    assert not rej
    for p in valid:
        qa = {}
        render_png(copy.deepcopy(p), tmp_path / f"{p['id']}.png", style=style, qa=qa)
        assert qa["spill"] == [], (p["id"], qa["spill"])


@pytest.mark.parametrize("style", ["classic", "enhanced"])
def test_the_emphasised_series_keeps_its_value_labels(tmp_path, style):
    pytest.importorskip("matplotlib")
    labels = [f"20{i:02d}" for i in range(18, 25)]
    data = {"years": labels}
    for i in range(3):
        data[f"s{i}"] = _series([12345.6 + 100 * i + 37 * j for j in range(7)], "億円",
                                kind="time", labels_ref="years")
    p = {"id": "x", "form": "bar", "title": "t", "labels_ref": "years",
         "series": [{"name": "s0", "values_ref": "s0", "role": "context"},
                    {"name": "s1", "values_ref": "s1", "role": "context2"},
                    {"name": "s2", "values_ref": "s2", "role": "emphasis"}],
         "caption": "c", "judgment": "j"}
    valid, _ = validate_panels([p], data)
    if style == "enhanced":
        valid[0]["form"] = "bar"                       # keep it vertical (7 labels would turn it)
    seen = _tick_texts(valid[0], style, tmp_path)
    emphasised = {f"{12345.6 + 200 + 37 * j:,.1f}" for j in range(7)}
    shown = {t.strip() for t in seen["texts"]}
    assert emphasised <= shown, emphasised - shown


@pytest.mark.parametrize("style", ["classic", "enhanced"])
def test_a_long_caption_and_legend_never_squeeze_the_plot_away(tmp_path, style):
    pytest.importorskip("matplotlib")
    data = {"years": ["2020", "2021", "2022"]}
    for i in range(8):
        data[f"s{i}"] = _series([10 + i, 11 + i, 12 + i], "億円", kind="time", labels_ref="years")
    p = {"id": "x", "form": "line_xy", "title": "t", "labels_ref": "years",
         "series": [{"name": f"とても長い系列の名前その{i}", "values_ref": f"s{i}",
                     "role": "emphasis" if i == 0 else "context"} for i in range(8)],
         "caption": "出所と定義の注記。" * 30, "judgment": "見るべき点の説明。" * 12}
    valid, _ = validate_panels([p], data)
    qa = {}
    render_png(valid[0], tmp_path / "x.png", style=style, qa=qa)
    assert qa["spill"] == [] and qa["label_overlaps"] == 0
    # the figure grows instead: at least the minimum plot height is left
    from deep_research_tool.chartkit.render_png import _LAYOUT
    assert qa["plot_in"] >= _LAYOUT[style]["min_plot"] - 0.6, qa     # (ticks / a wrapped title may take a little)


class TestReportRegressions:

    def _gen(self, tmp_path, coll):
        from deep_research_tool.report.generator import ReportGenerator
        from deep_research_tool.tests.test_finalization_integration import (
            make_locker, make_plan, make_session)
        ev = [("https://ex.com/a", "x " * 30)]
        plan = make_plan([("1", "市場規模の推移", "d")])
        session = make_session(plan, {"1": {"title": "市場規模の推移", "content": "本文 [1]", "evidence": ev}})
        gen = ReportGenerator(output_dir=tmp_path / "reports", language="ja")
        gen.figure_collection = coll
        return gen, session, make_locker(tmp_path, ev)

    def test_pdf_tables_with_long_cells_zero_values_and_no_title(self, tmp_path):
        pypdf = pytest.importorskip("pypdf")
        pytest.importorskip("reportlab")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.figure_table_generator import FigureTableCollection, TableData
        long_cell = "長い説明文。" * 400                     # taller than a page
        coll = FigureTableCollection(tables=[
            TableData(table_id="t1", title=None, caption="", headers=["指標", 0],
                      rows=[["A", 0], ["B", 0.0]], section_id="1"),
            TableData(table_id="t2", title="長い表", caption="", headers=["項目", "説明"],
                      rows=[["X", "短い"], ["Y", "中くらいの長さの説明" * 12]], section_id="1"),
            TableData(table_id="t3", title="巨大なセル", caption="", headers=["項目", "説明"],
                      rows=[["Z", long_cell]], section_id="1")])
        gen, session, locker = self._gen(tmp_path, coll)
        path = gen.generate_report(session=session, evidence_locker=locker,
                                   format=ReportFormat.PDF, verification_result=None)
        text = "".join(p.extract_text() for p in pypdf.PdfReader(str(path)).pages)
        assert gen.figures_inserted == 3
        compact = text.replace("\n", "").replace(" ", "")
        assert "A0" in compact and "B0.0" in compact               # zero cells are values
        assert compact.count("長い説明文。") >= 300                  # the giant cell is all there

    def test_html_text_with_a_less_than_sign_passes_the_artifact_check(self, tmp_path):
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.figure_table_generator import (
            Figure, FigureTableCollection, FigureType)
        from deep_research_tool.report.semantic_manifest import build_semantic_manifest, check_artifact
        png = tmp_path / "c.png"
        from PIL import Image
        Image.new("RGB", (400, 250), "white").save(png)
        coll = FigureTableCollection(charts=[Figure(
            figure_id="c1", figure_type=FigureType.CHART, title="伸び率",
            caption="伸び率が<1%の年は2回だけで、ほかは3%以上の伸びを続けた。", section_id="1",
            image_path=png)])
        gen, session, locker = self._gen(tmp_path, coll)
        session.section_contents["1"]["content"] = "本文 [1]\n\n伸び率が<5%の年は少なく、全体として堅調に推移したと言える。"
        path = gen.generate_report(session=session, evidence_locker=locker,
                                   format=ReportFormat.HTML, verification_result=None)
        manifest = build_semantic_manifest({"1": session.section_contents["1"]["content"]},
                                           figure_collection=coll)
        res = check_artifact(manifest, path, locker)
        assert res["status"] == "pass", res

    def test_v2_converter_renders_an_image_line(self):
        from deep_research_tool.report.v2.generator import ReportGeneratorV2
        html_text = ReportGeneratorV2._markdown_to_html("![図 \"1\"](/tmp/a b.png)\n\n![t](/tmp/c.png)")
        assert '<img src="/tmp/c.png" alt="t"' in html_text


def test_chart_warnings_only_for_rendered_charts(tmp_path, monkeypatch):
    pytest.importorskip("matplotlib")
    from types import SimpleNamespace
    from deep_research_tool.config import create_config
    from deep_research_tool.evidence.locker import EvidenceLocker
    from deep_research_tool.main import DeepResearchTool
    from deep_research_tool.utils.helpers import ResearchWarnings
    import deep_research_tool.chartkit as ck

    def boom(*a, **k):
        raise RuntimeError("renderer down")
    monkeypatch.setattr(ck, "render_png", boom)
    cfg = create_config(provider="openai", openai_api_key="t", output_dir=str(tmp_path), auto_figures=True)
    tool = DeepResearchTool.__new__(DeepResearchTool)
    tool.config, tool.llm_client, tool.stage_llm_clients = cfg, _LLM(), {}
    store = NumericalDataStore(research_topic="x")
    for y, v in ((2021, 40), (2022, 50), (2023, 55)):
        store.add(_dp("市場", "市場規模", y, v, "億円", "https://www.stat.go.jp/x", section_id="1",
                      raw_text="無関係な記述"))
    session = SimpleNamespace(query="x", research_plan=None,
                              section_contents={"1": {"title": "市場規模", "content": "本文", "sources": []}})
    collector = ResearchWarnings.bind()
    try:
        figs = tool._spec_charts(session, EvidenceLocker(research_id="r", output_dir=tmp_path / "ev"),
                                 store, tmp_path / "figures")
        msgs = [w["message"] for w in collector.to_dict_list()]
    finally:
        ResearchWarnings.unbind()
    assert figs == []
    assert not any("差し替え" in m or "図は表示しています" in m for m in msgs), msgs
    assert any("描画に失敗" in m for m in msgs), msgs


@pytest.mark.parametrize("version", ["v1", "v2"])
@pytest.mark.parametrize("fmt", ["markdown", "docx", "pdf", "html"])
def test_run_places_figures_in_every_format(tmp_path, version, fmt):
    """run() with a frozen figure collection: the report carries the chart
    (or the warning says why not) and the artifact check passes."""
    pytest.importorskip("reportlab")
    from unittest.mock import patch
    from PIL import Image
    from deep_research_tool.report.figure_table_generator import (
        Figure, FigureTableCollection, FigureTableGenerator, FigureType, TableData)
    from deep_research_tool.report.semantic_manifest import extract_artifact_text
    from deep_research_tool.tests import test_finalization_integration as T
    from deep_research_tool.utils.helpers import ResearchWarnings
    ResearchWarnings.reset()
    png = tmp_path / "figs" / "chart.png"
    png.parent.mkdir(parents=True)
    Image.new("RGB", (800, 500), "white").save(png)
    coll = FigureTableCollection(
        charts=[Figure(figure_id="c1", figure_type=FigureType.CHART, title="シェアの推移",
                       caption="A社のシェアの推移を示す図。出所: 統計", section_id="1", image_path=png)],
        tables=[TableData(table_id="t1", title="主要企業のシェア", caption="出所: 各社IR",
                          headers=["社名", "シェア"], rows=[["A社", "52%"], ["B社", "0"]], section_id="1")])
    llm = T.make_llm()
    if version == "v2":
        llm.on("chapter", lambda p: T.long_body("1", "市場シェア", T.CLAIM_SENTENCE)
               + "\n===CHAPTER_META===\n" + T.j({"key_points": [], "terms_used": [], "facts_stated": []}))
    locker = T.make_locker(tmp_path, T.SUPPORTED_EVIDENCE)
    tool = T.build_tool(tmp_path, llm, T.FakeSearch(), auto_figures=True, output_format=fmt,
                        report_generator_version=version)
    gen = FigureTableGenerator(llm_client=llm, output_dir=tmp_path / "figs")
    base = T.Base()
    base.tmp = tmp_path
    with patch.object(type(tool), "_generate_figure_collection", lambda self, **kw: (coll, gen)):
        res = T.run_tool(tool, base.default_session(), locker)
    assert res["semantic_artifact_check"] == "pass", res.get("semantic_artifact_check")
    path = Path(res["report_path"])
    if path.suffix == ".html":
        assert 'src="data:image/png;base64,' in path.read_text(encoding="utf-8")
    elif path.suffix == ".pdf":
        import pypdf
        assert sum(len(p.images) for p in pypdf.PdfReader(str(path)).pages) >= 1
    elif path.suffix == ".docx":
        import docx
        assert len(docx.Document(str(path)).inline_shapes) >= 1
    else:
        assert "chart.png" in path.read_text(encoding="utf-8")
    assert "A社のシェアの推移を示す図" in (extract_artifact_text(path) or "")


# --------------------------------------------------------------------------
# regressions from the second review round
# --------------------------------------------------------------------------

def _panel(form, labels, series, title="タイトル（億円）", unit="億円", **kw):
    """A resolved panel from plain values (labels_ref 'L')."""
    data = {"L": list(labels)}
    spec = []
    for i, (vals, role) in enumerate(series):
        key = f"s{i}"
        kind = "time" if form == "line_xy" else "comparison"
        data[key] = _series(list(vals), unit, kind=kind, labels_ref="L")
        spec.append({"name": kw.pop(f"name{i}", key), "values_ref": key, "role": role})
    p = {"id": "x", "form": form, "title": title, "labels_ref": "L", "series": spec,
         "caption": kw.pop("caption", "出所: テスト"), "judgment": kw.pop("judgment", "比較して見る。")}
    p.update(kw)
    valid, rej = validate_panels([p], data)
    assert not rej, rej
    return valid[0]


def _qa(panel, style, tmp_path, name="q"):
    qa = {}
    render_png(copy.deepcopy(panel), tmp_path / f"{name}_{style}.png", style=style, qa=qa)
    return qa


@pytest.mark.parametrize("style", ["classic", "enhanced"])
def test_r2_layout_cases_stay_readable(tmp_path, style):
    pytest.importorskip("matplotlib")
    years = [str(y) for y in range(2019, 2024)]
    cases = {
        # a long title is wrapped, not cut at the image edge
        "long_title": _panel("bar", ["A", "B", "C"], [([3, 2, 1], "emphasis")],
                             title="国内ノンアルコールビール市場の販売数量推移（メーカー別、2019〜2024年）の比較と今後の見通し（億円）"),
        # thousands-separated ticks on a horizontal axis do not print over each other
        "hbar_big": _panel("hbar", ["A", "B", "C", "D"], [([125e6, 98e6, 64e6, 12e6], "emphasis")], unit="人"),
        # a negative bar's label never prints over the category name
        "hbar_neg": _panel("hbar", ["営業利益", "経常利益", "純利益"], [([-21, 32, -15], "emphasis")]),
        # the first point is the maximum: its label is not printed on the title
        "line_top": _panel("line_xy", years, [([160, 120, 110, 100, 90], "emphasis")]),
        # 20 categories in a stacked bar: the tick labels never overlap
        "stack20": _panel("stacked_bar", [f"カテゴリ{i}" for i in range(20)],
                          [([i + 1 for i in range(20)], "part"), ([20 - i for i in range(20)], "part")]),
        # a series name wider than the image wraps in the legend
        "long_legend": _panel("line_xy", years, [([1, 2, 3, 4, 5], "emphasis"), ([2, 3, 4, 5, 6], "context")],
                              name0="とても長い系列の名前" * 10),
    }
    for name, panel in cases.items():
        qa = _qa(panel, style, tmp_path, name)
        assert qa["spill"] == [], (name, qa["spill"])
        assert qa["tick_overlaps"] == 0 and qa["label_overlaps"] == 0, (name, qa)
        assert qa["label_on_title"] == 0 and qa["label_on_ticks"] == 0, (name, qa)
        if style == "enhanced":
            assert qa["size_px"][0] == 1100, (name, qa["size_px"])     # still printed 1:1


def test_r2_tiny_values_get_real_tick_text(tmp_path):
    pytest.importorskip("matplotlib")
    panel = _panel("bar", ["A", "B", "C"], [([1e-6, 2.5e-6, 4e-6], "emphasis")], unit="円")
    ticks = _tick_texts(panel, "enhanced", tmp_path)["ticks"]
    assert len({t for _l, t in ticks}) == len(ticks) and any(t != "0" for _l, t in ticks)


def test_r2_enhanced_keeps_several_series_vertical(tmp_path):
    pytest.importorskip("matplotlib")
    panel = _panel("bar", [f"P{i}" for i in range(8)],
                   [([i + 1 for i in range(8)], "emphasis"), ([i + 2 for i in range(8)], "context")])
    render_png(panel, tmp_path / "m.png", style="enhanced")
    assert "form_rendered" not in panel


def test_r2_the_emphasised_line_keeps_its_point_labels(tmp_path):
    pytest.importorskip("matplotlib")
    years = [str(y) for y in range(2019, 2024)]
    panel = _panel("line_xy", years, [([10, 11, 12, 13, 14], "context"), ([10.2, 11.2, 12.2, 13.2, 14.2], "emphasis")],
                   unit="%")
    shown = {t.strip() for t in _tick_texts(panel, "classic", tmp_path)["texts"]}
    assert {"10.2%", "11.2%", "12.2%", "13.2%", "14.2%"} <= shown


class TestR2ClassicIsV15:
    """classic draws what v1.5 drew wherever v1.5 had no defect."""

    def test_fonts_and_heights(self, tmp_path):
        pytest.importorskip("matplotlib")
        from deep_research_tool.chartkit.render_png import _JP_FONT_HINTS
        assert _JP_FONT_HINTS["classic"][0] == "IPAGothic"            # the v1.5 preference order
        hbar = _panel("hbar", list("ABCDE"), [([5, 4, 3, 2, 1], "emphasis")])
        assert _qa(hbar, "classic", tmp_path)["size_px"] == (1500, round(max(3.2, 0.45 * 5 + 2.0) * 150))
        seg = _panel("stacked_hbar", ["X"], [([60], "part"), ([40], "part")], unit="%")
        assert _qa(seg, "classic", tmp_path, "s")["size_px"] == (1500, round(max(2.6, 0.55 + 2.0) * 150))

    def test_value_axis_is_not_extended_for_a_label_inside_the_image(self, tmp_path):
        pytest.importorskip("matplotlib")
        hbar = _panel("hbar", list("ABCDE"), [([12345, 9876, 7500, 5100, 2300], "emphasis")], unit="百万円")
        lo, hi = _qa(hbar, "classic", tmp_path)["value_lim"]
        assert hi == pytest.approx(12345 * 1.05, rel=1e-3)           # matplotlib's own margin, as in v1.5

    def test_judgment_and_labels(self, tmp_path):
        pytest.importorskip("matplotlib")
        judgment = ("2024年度の売上高は前年比12.5%増の3,456億円、営業利益率は8.2%で、"
                    "2020年度以降の最高水準を5年連続で更新した。国内好調")
        panel = _panel("hbar", list("ABC"), [([3, 2, 1], "emphasis")], judgment=judgment)
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        seen = {}
        orig = plt.Figure.savefig

        def spy(self, *a, **k):
            seen.setdefault("fig", []).extend(t.get_text() for t in self.texts)
            seen["va"] = [t.get_va() for t in self.axes[0].texts]     # the line chart (last)
            return orig(self, *a, **k)
        plt.Figure.savefig = spy
        try:
            render_png(copy.deepcopy(panel), tmp_path / "j.png", style="classic")
            line = _panel("line_xy", ["2020", "2021"], [([1, 2], "emphasis")])
            render_png(copy.deepcopy(line), tmp_path / "l.png", style="classic")
        finally:
            plt.Figure.savefig = orig
        assert any(t == "見るべき点: " + judgment for t in seen["fig"])     # one line, as in v1.5
        assert set(seen["va"]) == {"baseline"}                              # point labels as in v1.5


class TestR2Data:

    def test_scale_words(self):
        from deep_research_tool.chartkit.data_layer import _raw_numbers, parse_unit
        assert repair_scale(5, "5 tn yen", "billion") == ("trillion", True)
        assert parse_unit(repair_scale(5, "5 tn yen", "billion")[0])[1] == 1e12
        assert [w for _n, _f, w in _raw_numbers("BNPL 5 MNOs 3 billionaires")] == ["", ""]

    def test_japanese_compound_numbers_and_wari(self):
        for v, unit, quote in [(3.5, "億円", "3億5,000万円"), (12000, "億円", "1兆2,000億円"),
                               (1.2, "兆円", "1兆2000億円"), (1.2, "億人", "1億2,000万人"),
                               (0.3, "億円", "3千万円"), (1000, "億円", "1千億円"), (1.5, "万円", "1万5000円"),
                               (30, "%", "約3割"), (35, "%", "3割5分")]:
            assert value_in_quote(v, unit, quote), (v, unit, quote)
        assert not value_in_quote(35, "億円", "3億5,000万円")              # still a 10x error

    def test_table_rows_do_not_fuse(self):
        page = "（単位：億円）\n2022年度  2023年度  2024年度\n売上高   1,100   1,200   1,350\n"
        assert quote_in_source("売上高 1,200", page) and quote_in_source("1,350", page)
        assert quote_in_source("23.5%", "Share 2022 21.0% 2023 23.5% 2024 24.1%")
        assert not quote_in_source("1,000", page)


class TestR2Reports:

    def _gen(self, tmp_path, coll, title=None):
        from deep_research_tool.report.generator import ReportGenerator
        from deep_research_tool.tests.test_finalization_integration import (
            make_locker, make_plan, make_session)
        ev = [("https://ex.com/a", "x " * 30)]
        plan = make_plan([("1", "市場規模の推移", "d")])
        if title:
            plan.title = title
        session = make_session(plan, {"1": {"title": "市場規模の推移", "content": "本文 [1]", "evidence": ev}})
        gen = ReportGenerator(output_dir=tmp_path / "reports", language="ja")
        gen.figure_collection = coll
        return gen, session, make_locker(tmp_path, ev)

    def test_pdf_survives_wide_tables_and_cut_off_images(self, tmp_path):
        pytest.importorskip("reportlab")
        from PIL import Image
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.figure_table_generator import (
            Figure, FigureTableCollection, FigureType, TableData)
        good = tmp_path / "good.png"
        Image.new("RGB", (600, 400), "white").save(good)
        cut = tmp_path / "cut.png"
        Image.new("RGB", (600, 400), (10, 120, 200)).save(cut)
        cut.write_bytes(cut.read_bytes()[: len(cut.read_bytes()) // 2])
        coll = FigureTableCollection(
            figures=[Figure(figure_id="f1", figure_type=FigureType.IMAGE, title="切れた画像", caption="c1",
                            section_id="1", image_path=cut),
                     Figure(figure_id="f2", figure_type=FigureType.IMAGE, title="正常な画像", caption="c2",
                            section_id="1", image_path=good)],
            tables=[TableData(table_id="w", title="40列", caption="", headers=[f"{2000 + i}" for i in range(40)],
                              rows=[[str(i) for i in range(40)]], section_id="1"),
                    TableData(table_id="l", title="長いセル", caption="", headers=[f"h{i}" for i in range(12)],
                              rows=[["長い説明" * 35] * 12], section_id="1")])
        gen, session, locker = self._gen(tmp_path, coll)
        path = gen.generate_report(session=session, evidence_locker=locker,
                                   format=ReportFormat.PDF, verification_result=None)
        assert path.exists() and gen.figures_inserted == 3                # the cut image is skipped

    @pytest.mark.parametrize("fmt", ["html", "pdf", "markdown"])
    def test_injected_text_is_caught_and_titles_are_allowed(self, tmp_path, fmt):
        pytest.importorskip("reportlab")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.figure_table_generator import FigureTableCollection, TableData
        from deep_research_tool.report.semantic_manifest import build_semantic_manifest, check_artifact
        from deep_research_tool.main import _report_titles
        coll = FigureTableCollection(tables=[TableData(
            table_id="t", title="主要企業", caption="", headers=["社名", "シェア"],
            rows=[["A社", "52%"], ["B社", "0"]], section_id="1")])
        long_title = "日本国内における炭素繊維強化プラスチック市場の動向と主要企業の競争戦略および今後の成長見通しに関する調査報告"
        gen, session, locker = self._gen(tmp_path, coll, title=long_title)
        body = ("本文 [1]\n\n国内市場は堅調に拡大しており、主要メーカーは増産投資を進めている。"
                "一方で原料価格の上昇が収益を圧迫する懸念も残る。")
        session.section_contents["1"]["content"] = body
        # Markdown gets its figures after rendering (FigureTableGenerator),
        # PDF / HTML place them while rendering
        if fmt == "markdown":
            gen.figure_collection = coll = None
        manifest = build_semantic_manifest({"1": body}, figure_collection=coll)
        path = gen.generate_report(session=session, evidence_locker=locker,
                                   format=ReportFormat(fmt), verification_result=None)
        assert check_artifact(manifest, path, locker,
                              extra_allowed=_report_titles(session))["status"] == "pass"
        # the same report with one sentence added after the freeze fails
        injected = "競合他社の売上高は2030年までに三倍に急増し今後も高成長が確実視されている状況だと断言できる。"
        session.section_contents["1"]["content"] = body + "\n\n" + injected
        path2 = gen.generate_report(session=session, evidence_locker=locker,
                                    format=ReportFormat(fmt), verification_result=None)
        res = check_artifact(manifest, path2, locker, extra_allowed=_report_titles(session))
        assert res["status"] == "fail" and res["additions"], res

    def test_missing_image_comment_is_not_text(self):
        from deep_research_tool.report.semantic_manifest import normalize_semantic_text
        assert normalize_semantic_text("<!-- Chart not found: /tmp/x/chart_s1_p1.png -->") == ""

    def test_v2_html_tables_and_odd_image_paths(self):
        from deep_research_tool.report.v2.generator import ReportGeneratorV2
        page = ReportGeneratorV2._markdown_to_html(
            "| 社名 | シェア |\n|---|---|\n| A社 | 52% |\n\n![題 [注]](/tmp/My Reports/c.png)")
        assert "<td>A社</td><td>52%</td>" in page
        assert '<img src="/tmp/My Reports/c.png" alt="題 [注]"' in page


# --------------------------------------------------------------------------
# regressions from the third review round
# --------------------------------------------------------------------------

def test_r3_wrap_matches_v15_textwrap():
    import textwrap
    from deep_research_tool.chartkit.render_png import _wrap
    for text, width in [("生産能力 2026年末", 10), ("2026 出荷 年率換算", 10), ("日本 市場規模 2025年", 10),
                        ("Global market size", 10), ("出典: " + "あ" * 79, 70),
                        ("出典：…「工業統計」　各社決算資料　（2024年）", 70), ("a  b   c", 10)]:
        assert _wrap(text, width) == "\n".join(textwrap.wrap(text, width)), text
    # where textwrap had a defect the wrap differs on purpose
    assert "1.\n6" not in _wrap("26Q1の実売ベースは1Q'25基準1.6億缶×実売+10%からの推定。" * 2, 70)


def test_r3_segment_label_that_fits_keeps_its_name(tmp_path):
    pytest.importorskip("matplotlib")
    panel = _panel("stacked_hbar", ["2022年", "2023年", "2024年"],
                   [([55, 50, 46], "part"), ([30, 32, 35], "part"), ([10, 12, 13], "part"), ([5, 6, 6], "part")],
                   unit="%", y_max=100, name0="鉄鋼", name1="化学", name2="機械", name3="食品")
    texts = {t.strip() for t in _tick_texts(panel, "classic", tmp_path)["texts"]}
    assert "鉄鋼 55%" in texts and "機械 10%" in texts


def test_r3_enhanced_multi_series_hbar_fits_a_page(tmp_path):
    pytest.importorskip("matplotlib")
    panel = _panel("hbar", [f"C{i}" for i in range(10)],
                   [([i + j for i in range(10)], "emphasis" if j == 0 else "context") for j in range(8)])
    qa = _qa(panel, "enhanced", tmp_path)
    assert qa["size_px"][1] <= 10.5 * 200 and qa["spill"] == []


class TestR3Facts:

    def _line(self, vals, unit="億円", start=2019, forecast=()):
        return {"form": "line_xy", "unit": unit, "labels": [], "title": "t", "caption": "", "judgment": "",
                "series": [{"name": "s", "unit": unit, "basis": [PUB] * len(vals),
                            "values": [{"x": float(start + i), "y": v} for i, v in enumerate(vals)],
                            "sources": [{"forecast": (start + i) in forecast} for i in range(len(vals))]}]}

    def test_true_sentences_are_never_flagged(self):
        line = panel_facts(self._line([100, 72, 85, 95, 104]))
        for t in ("2020年は最も大きく落ち込んだ", "2021年は最も大きく伸びた", "2021年は過去最高を更新できなかった",
                  "2023年は過去最高だった2019年の約8割まで回復した", "2020年は過去最大となる落ち込みを記録した",
                  "市場は毎年拡大を続けている", "原材料費が毎年上昇し、利益率を圧迫している",
                  "需要はコロナ禍後、一貫して回復している"):
            assert judgment_conflict(t, line) is None, t
        bar = panel_facts({"form": "bar", "unit": "億円", "labels": ["A社", "B社", "C社"],
                           "series": [{"values": [320, 300, 150]}]})
        for t in ("B社は首位を奪われ、2位に後退した", "B社はトップを維持できなかった", "B社は首位に立てなかった"):
            assert judgment_conflict(t, bar) is None, t
        city = panel_facts({"form": "bar", "unit": "万人泊", "labels": ["東京", "京都"],
                            "series": [{"values": [5800, 1900]}]})
        assert judgment_conflict("延べ宿泊者数は東京都が最多", city) is None
        rank = panel_facts({"form": "bar", "unit": "位", "labels": ["シンガポール", "日本"],
                            "series": [{"values": [1, 38]}]})
        assert judgment_conflict("シンガポールが1位", rank) is None

    def test_data_sentence_numbers_and_units(self):
        from deep_research_tool.chartkit.facts import fact_sentence
        idx = panel_facts(self._line([98.5, 102, 105, 108.9], unit="2015年=100", start=2020))
        assert "98.5（2015年=100）→108.9（2015年=100）" in fact_sentence(idx)
        loss = panel_facts(self._line([-120, -90, -60, -20], start=2020))
        assert "%" not in fact_sentence(loss).split("（")[-1]           # no '+83%' across a negative base
        fc = panel_facts(self._line([100, 150, 200, 260], start=2022, forecast=(2025,)))
        assert "2022→2025（予測）" in fact_sentence(fc)
        rank = panel_facts({"form": "bar", "unit": "位", "labels": ["シンガポール", "日本"],
                            "series": [{"values": [1, 38]}]})
        assert fact_sentence(rank).startswith("最上位はシンガポール（1位）")


class TestR3Grounding:

    def test_scale_errors_are_flagged(self):
        assert not value_in_quote(3, "%", "シェアは約3割")                 # 30 %, not 3 %
        assert not value_in_quote(5, "%", "3割5分")
        assert not value_in_quote(5.2, "億ドル", "$5.2B")                  # 52 億ドル
        assert not value_in_quote(450, "億ドル", "US$450M")
        assert not value_in_quote(50, "%", "利用者は50万人")
        assert not value_in_quote(1234, "億円", "1,243億円")

    def test_more_written_forms_are_read(self):
        assert value_in_quote(30, "%", "約三割") and value_in_quote(35, "%", "3割5分")
        assert value_in_quote(33.3, "%", "3分の1") and value_in_quote(25, "%", "4人に1人")
        assert value_in_quote(12, "億ドル", "$1.2B") and value_in_quote(4.5, "億ドル", "US$450M")
        assert value_in_quote(1.5, "千", "1.5k units")
        assert quote_in_source("12.5%", "伸び率は12.50%") and quote_in_source("1.2兆円", "市場は1兆2,000億円")

    def test_compound_units_parse(self):
        from deep_research_tool.chartkit.data_layer import parse_unit
        assert parse_unit("千億円")[:2] == ("円", 1e11) and parse_unit("100万台")[:2] == ("台", 1e6)
        assert value_in_quote(3, "千億円", "2030年に3千億円") and value_in_quote(1.2, "100万台", "120万台")


class TestR3Artifact:

    def _gen(self, tmp_path, coll=None, plan_items=(("1", "市場規模の推移", "d"),)):
        from deep_research_tool.report.generator import ReportGenerator
        from deep_research_tool.tests.test_finalization_integration import (
            make_locker, make_plan, make_session)
        ev = [("https://ex.com/a", "x " * 30)]
        plan = make_plan(list(plan_items))
        session = make_session(plan, {sid: {"title": t, "content": "本文 [1]", "evidence": ev}
                                      for sid, t, _d in plan_items})
        gen = ReportGenerator(output_dir=tmp_path / "reports", language="ja")
        gen.figure_collection = coll
        return gen, session, make_locker(tmp_path, ev)

    def test_wide_table_keeps_its_grid_and_passes(self, tmp_path):
        pytest.importorskip("reportlab")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.figure_table_generator import FigureTableCollection, TableData
        from deep_research_tool.report.semantic_manifest import (build_semantic_manifest, check_artifact,
                                                                 figure_semantics_markdown)
        from deep_research_tool.main import _report_titles
        coll = FigureTableCollection(tables=[TableData(
            table_id="w", title="年次推移", caption="", headers=["メーカー"] + [str(y) for y in range(2011, 2024)],
            rows=[["A社"] + [str(10 + i) for i in range(13)]], section_id="1")])
        gen, session, locker = self._gen(tmp_path, coll)
        app = figure_semantics_markdown(coll, language="ja", section_id="付録D")
        session.section_contents["付録D"] = {"title": "付録D", "content": app, "sources": [], "extracted_content": []}
        path = gen.generate_report(session=session, evidence_locker=locker, format=ReportFormat.PDF,
                                   verification_result=None)
        manifest = build_semantic_manifest({"1": "本文 [1]", "付録D": app}, figure_collection=coll)
        assert check_artifact(manifest, path, locker, extra_allowed=_report_titles(session))["status"] == "pass"

    @pytest.mark.parametrize("fmt", ["markdown", "html", "pdf"])
    def test_long_chapter_titles_and_reference_prefixes(self, tmp_path, fmt):
        pytest.importorskip("reportlab")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.semantic_manifest import build_semantic_manifest, check_artifact
        from deep_research_tool.main import _report_titles
        long_title = "国内炭素繊維市場における主要メーカーの競争環境と市場シェア分析および今後の見通し"
        gen, session, locker = self._gen(tmp_path, plan_items=(("1", long_title, "d"),))
        body = "本文 [1]\n\n国内市場は堅調に拡大しており、主要メーカーは増産投資を進めている状況にある。"
        session.section_contents["1"]["content"] = body
        manifest = build_semantic_manifest({"1": body})
        path = gen.generate_report(session=session, evidence_locker=locker, format=ReportFormat(fmt),
                                   verification_result=None)
        extra = _report_titles(session)
        assert check_artifact(manifest, path, locker, extra_allowed=extra)["status"] == "pass"
        for injected in ("[9]: 競合他社の売上高は2030年までに三倍に急増し、今後も高い成長が確実視されていると経営陣が表明した。",
                         "1. [注] 競合他社の売上高は2030年までに三倍に急増し、今後も高い成長が確実視されていると経営陣が表明した。"):
            session.section_contents["1"]["content"] = body + "\n\n" + injected
            p2 = gen.generate_report(session=session, evidence_locker=locker, format=ReportFormat(fmt),
                                     verification_result=None)
            assert check_artifact(manifest, p2, locker, extra_allowed=extra)["status"] == "fail", injected

    def test_rules_bullets_and_numbered_lists_normalise(self):
        from deep_research_tool.report.semantic_manifest import normalize_semantic_text, verify_frozen_in_artifact
        assert normalize_semantic_text("─" * 60) == "" and normalize_semantic_text("\x7f 本文") == "本文"
        manifest = {"chapters": {"1": "1. 国内市場は堅調に拡大しており主要メーカーは増産投資を進めている。"}}
        res = verify_frozen_in_artifact(manifest, "国内市場は堅調に拡大しており主要メーカーは増産投資を進めている。")
        assert res["ok"], res

    def test_pdf_subsection_is_not_cut(self, tmp_path):
        pytest.importorskip("reportlab")
        pypdf = pytest.importorskip("pypdf")
        from deep_research_tool.config import ReportFormat
        gen, session, locker = self._gen(tmp_path, plan_items=(("1", "市場", "d"), ("1.1", "詳細", "d")))
        session.research_plan.table_of_contents.items[0].subsections = [
            i for i in session.research_plan.table_of_contents.get_flat_sections() if i.section == "1.1"]
        session.research_plan.table_of_contents.items = session.research_plan.table_of_contents.items[:1]
        long_text = "詳細な説明が続く。" * 300 + "最後の文はここにある。"
        session.section_contents["1.1"]["content"] = long_text
        path = gen.generate_report(session=session, evidence_locker=locker, format=ReportFormat.PDF,
                                   verification_result=None)
        text = "".join(p.extract_text() for p in pypdf.PdfReader(str(path)).pages).replace("\n", "")
        assert "最後の文はここにある。" in text


# --------------------------------------------------------------------------
# regressions from the fourth review round
# --------------------------------------------------------------------------

def test_r4_wrap_matches_textwrap_on_hyphens_tabs_and_indents():
    import textwrap
    from deep_research_tool.chartkit.render_png import _wrap
    for text, width in [("Fuji-Xerox Holdings market share", 10), ("well-known brand-new state-of-the-art", 12),
                        ("  indented text here", 10), ("a\tb c", 10), ("line1\r\nline2", 10)]:
        assert _wrap(text, width) == "\n".join(textwrap.wrap(text, width)), text
    # a number keeps its 億 / 万 multiplier at a line end
    for width in range(14, 20):
        assert "6\n億" not in _wrap("あいうえおかきくけこさしすせそ6億缶の出荷", width)


@pytest.mark.parametrize("n_cats,n_series", [(15, 3), (36, 1)])
def test_r4_enhanced_hbar_height_is_bounded(tmp_path, n_cats, n_series):
    pytest.importorskip("matplotlib")
    panel = _panel("hbar", [f"カテゴリー{i}" for i in range(n_cats)],
                   [([10 + i + j for i in range(n_cats)], "emphasis" if j == 0 else "context")
                    for j in range(n_series)])
    qa = _qa(panel, "enhanced", tmp_path)
    assert qa["size_px"][1] <= 10.5 * 200 and qa["spill"] == [] and not qa["tick_overlaps"]


class TestR4Units:

    def test_scale_letters(self):
        from deep_research_tool.chartkit.data_layer import _raw_numbers, repair_scale
        assert _raw_numbers("10kΩ") == [(10.0, 1.0, "")]                 # a unit symbol, not x1000
        assert repair_scale(10, "抵抗値は10kΩ", "kΩ") == ("kΩ", False)
        assert _raw_numbers("US$450m")[0][1] == 1e6 and _raw_numbers("USD 1.2B")[0][1] == 1e9
        assert _raw_numbers("$1.2T")[0][1] == 1e12 and _raw_numbers("5m")[0][1] == 1.0
        assert _raw_numbers("3百億円")[0][2] == "?"                        # compound: not a bare 3

    def test_compound_and_rate_units(self):
        from deep_research_tool.chartkit.data_layer import parse_unit, repair_scale
        assert parse_unit("1000億円")[:2] == ("円", 1e11) and parse_unit("1,000万円")[:2] == ("円", 1e7)
        # a per-population rate is a unit, not a scale
        assert parse_unit("人口10万人当たり")[1] == 1.0 and parse_unit("百万人当たり")[1] == 1.0
        assert repair_scale(5.1, "revenue of $5.1 billion", "千億円") == ("billion 円", True)

    def test_scale_errors_in_grounding(self):
        assert not value_in_quote(450, "億ドル", "US$450m") and value_in_quote(4.5, "億ドル", "US$450m")
        assert not value_in_quote(1.2, "億ドル", "$1.2T") and value_in_quote(1.2, "兆ドル", "$1.2T")
        assert not value_in_quote(1.2, "億ドル", "USD 1.2B") and value_in_quote(12, "億ドル", "USD 1.2B")
        assert not value_in_quote(3, "億円", "3百億円") and value_in_quote(300, "億円", "3百億円")

    def test_quotes_against_the_page(self):
        from deep_research_tool.chartkit.grounding import _scale_conflict
        assert not quote_in_source("売上高は1,234億円", "売上高は1,240億円だった")   # a misquote
        assert quote_in_source("1兆2,000億円", "市場は1.2兆円") and quote_in_source("3億5,000万円", "3.5億円")
        assert not quote_in_source("$5.2 billion", "revenue was $5.2 million")
        assert _scale_conflict("$5.2 billion", "revenue was $5.2 million")

    def test_rounded_fractions(self):
        assert value_in_quote(33, "%", "3人に1人") and value_in_quote(33.3, "%", "3人に1人")
        assert value_in_quote(14, "%", "7人に1人") and value_in_quote(67, "%", "3分の2")


class TestR4Facts:

    def _line(self, vals, unit="億円", start=2020):
        return panel_facts({"form": "line_xy", "unit": unit, "labels": [], "title": "t", "caption": "",
                            "judgment": "", "series": [{"name": "s", "unit": unit, "basis": [PUB] * len(vals),
                                                        "values": [{"x": float(start + i), "y": v}
                                                                   for i, v in enumerate(vals)],
                                                        "sources": [{} for _ in vals]}]})

    def test_rank_and_small_percent_wording(self):
        from deep_research_tool.chartkit.facts import fact_sentence
        assert fact_sentence(self._line([12, 9, 5, 3], unit="位")) == "2020→2023で12位→3位（順位を上げた）"
        assert "増加、+0.5pt" in fact_sentence(self._line([0.5, 0.6, 0.8, 1.0], unit="%"))
        assert "+0.03pt" in fact_sentence(self._line([0.5, 0.53], unit="%"))

    def test_scope_and_case(self):
        bar = panel_facts({"form": "bar", "unit": "億円", "labels": ["toyota", "honda", "ｓｏｎｙ"],
                           "series": [{"values": [100, 300, 50]}]})
        assert judgment_conflict("Toyotaが首位。", bar) and judgment_conflict("ＳＯＮＹが首位。", bar)
        assert judgment_conflict("Hondaが首位。", bar) is None
        for t in ("国内ではToyotaが首位。", "国内市場に限ればToyotaが首位", "北米だけで見るとToyotaが首位"):
            assert judgment_conflict(t, bar) is None, t


class TestR4Reports:

    def test_links_with_quotes_and_parentheses(self):
        from deep_research_tool.report.generator import _html_linked, _pdf_links
        from deep_research_tool.report.semantic_manifest import normalize_semantic_text
        url = "https://ja.wikipedia.org/wiki/ノンアルコールビール_(日本)"
        assert f'<a href="{url}"' in _pdf_links(f"制度は[解説記事]({url})が詳しい。")
        assert normalize_semantic_text(f"制度は[解説記事]({url})が詳しい。") == "制度は解説記事が詳しい。"
        assert '%22' in _pdf_links('[記事](https://ex.com/a"b)') and '"b' not in _pdf_links('[記事](https://ex.com/a"b)')
        assert _html_linked("出所: [統計](https://ex.com/s?a=1&b=2)") == \
            '出所: <a href="https://ex.com/s?a=1&amp;b=2">統計</a>'

    @pytest.mark.parametrize("fmt", ["pdf", "html"])
    def test_links_in_captions_and_cells_print_as_text(self, tmp_path, fmt):
        pytest.importorskip("reportlab")
        from PIL import Image
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.figure_table_generator import (
            Figure, FigureTableCollection, FigureType, TableData)
        from deep_research_tool.report.semantic_manifest import extract_artifact_text
        png = tmp_path / "c.png"
        Image.new("RGB", (600, 400), "white").save(png)
        coll = FigureTableCollection(
            charts=[Figure(figure_id="c1", figure_type=FigureType.CHART, title="推移", section_id="1",
                           caption="シェアの推移。出所: [業界団体の出荷統計](https://www.example.or.jp/stats/2024/a.html)",
                           image_path=png)],
            tables=[TableData(table_id="t1", title="主要企業", section_id="1",
                              caption="出所: [商業動態統計](https://www.meti.go.jp/statistics/tyo/syoudou/x.xlsx)",
                              headers=["資料", "内容"], rows=[["[経済産業省](https://www.meti.go.jp/a.html)", "月次"]])])
        gen, session, locker = TestR3Artifact()._gen(tmp_path, coll)
        path = gen.generate_report(session=session, evidence_locker=locker, format=ReportFormat(fmt),
                                   verification_result=None)
        text = extract_artifact_text(path)
        assert "業界団体の出荷統計" in text and "商業動態統計" in text and "経済産業省" in text
        assert "](http" not in text

    def test_english_pdf_takes_a_cjk_font_for_a_japanese_summary(self, tmp_path, monkeypatch):
        pytest.importorskip("reportlab")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.generator import ReportGenerator
        from deep_research_tool.tests.test_finalization_integration import make_locker, make_plan, make_session
        calls = []
        monkeypatch.setattr(ReportGenerator, "_register_japanese_fonts", lambda self: calls.append(1))
        plan = make_plan([("1", "Market share", "d")])
        session = make_session(plan, {"1": {"title": "Market share", "content": "Company A leads. " * 5,
                                            "evidence": []}})
        session.section_contents["_executive_summary"] = {
            "executive_summary": "The market keeps growing.", "recommendations": [],
            "key_findings": ["The flagship brand 'オールフリー' leads the category nationwide"]}
        gen = ReportGenerator(output_dir=tmp_path, language="en", include_citations=False)
        gen.generate_report(session, make_locker(tmp_path, [("https://ex.com/a", "data " * 20)]),
                            format=ReportFormat.PDF, filename="r")
        assert calls

    def test_wide_table_numbers_stay_whole_in_the_pdf(self, tmp_path):
        pytest.importorskip("reportlab")
        pypdf = pytest.importorskip("pypdf")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.figure_table_generator import FigureTableCollection, TableData
        coll = FigureTableCollection(tables=[TableData(
            table_id="w", title="年次推移", caption="", headers=["メーカー"] + [str(y) for y in range(2011, 2024)],
            rows=[["A社"] + [f"12,{345 + i}.6" for i in range(13)]], section_id="1")])
        gen, session, locker = TestR3Artifact()._gen(tmp_path, coll)
        path = gen.generate_report(session=session, evidence_locker=locker, format=ReportFormat.PDF,
                                   verification_result=None)
        lines = "\n".join(p.extract_text() for p in pypdf.PdfReader(str(path)).pages).splitlines()
        text = " ".join(lines)
        assert "2011" in text and "2023" in text and "12,345.6" in text and "12,357.6" in text

    def test_numbered_fallback_takes_integer_list_numbers_only(self):
        from deep_research_tool.report.semantic_manifest import verify_frozen_in_artifact
        # a line under the addition floor and a table cell: only the missing check sees them
        frozen = "1.5兆円規模の国内市場は2030年までに倍増する見込みである。"
        res = verify_frozen_in_artifact({"chapters": {"1": frozen}}, frozen.replace("1.5兆円", "2.5兆円"))
        assert not res["ok"]
        body = "国内市場は2030年までに大きく拡大する見込みであり、主要各社は投資を続けている。"
        table = {"chapters": {"1": body}, "figures": [{"title": "表1 市場規模", "caption": "",
                                                       "headers": ["年", "成長率"], "rows": [["2024", "12.5%"]]}]}
        assert not verify_frozen_in_artifact(table, body + "\n表1 市場規模\n年 成長率\n2024 99.5%")["ok"]
        assert verify_frozen_in_artifact(table, body + "\n表1 市場規模\n年 成長率\n2024 12.5%")["ok"]
        listed = "1. 国内市場は堅調に拡大しており主要メーカーは増産投資を進めている。"
        assert verify_frozen_in_artifact({"chapters": {"1": listed}}, listed[3:])["ok"]

    def test_docx_tall_chart_fits_the_page(self, tmp_path):
        docx = pytest.importorskip("docx")
        from PIL import Image
        from docx.shared import Inches
        from deep_research_tool.report.figure_table_generator import _add_docx_picture
        tall, wide = tmp_path / "tall.png", tmp_path / "wide.png"
        Image.new("RGB", (400, 2000), "white").save(tall)
        Image.new("RGB", (1100, 700), "white").save(wide)
        doc = docx.Document()
        _add_docx_picture(doc.add_paragraph().add_run(), tall, Inches)
        _add_docx_picture(doc.add_paragraph().add_run(), wide, Inches)
        shapes = doc.inline_shapes
        assert shapes[0].height <= Inches(8.0) + 1 and shapes[1].width == Inches(5.5)

    def test_resume_shows_the_chart_warnings_again(self, tmp_path):
        from deep_research_tool.report.figure_table_generator import FigureTableCollection
        from deep_research_tool.tests.test_finalization_integration import FakeSearch, build_tool, make_llm
        from deep_research_tool.utils.helpers import ResearchWarnings
        figs = tmp_path / "figs"
        figs.mkdir()
        (figs / "figure_collection.json").write_text(json.dumps(FigureTableCollection().to_dict()), encoding="utf-8")
        msg = "図「市場規模」: 判断文が数値と食い違う可能性があります"
        (figs / "chart_yield.json").write_text(json.dumps({"warnings": [msg]}, ensure_ascii=False), encoding="utf-8")
        ResearchWarnings.reset()
        tool = build_tool(tmp_path, make_llm(), FakeSearch())
        coll, _gen = tool._restore_figures(figs / "figure_collection.json")
        assert coll is not None
        assert any(msg in w["message"] for w in ResearchWarnings.get_instance().get_all())


# --------------------------------------------------------------------------
# regressions from the fifth review round
# --------------------------------------------------------------------------

def test_r5_long_hyphenated_words_break_at_the_hyphen_as_textwrap():
    import textwrap
    from deep_research_tool.chartkit.render_png import _wrap
    for text, width in [("FY2020-FY2025", 10), ("gpt-4o-mini", 10), ("claude-3-opus", 10), ("gemini-1.5-pro", 10),
                        ("e-commerce", 8), ("2020-2025年度", 8), ("abcdefghi claude-3-opus x", 10),
                        ("  Revenue growth", 8), ("x-ray imaging", 6)]:
        assert _wrap(text, width) == "\n".join(textwrap.wrap(text, width)), text
    assert _wrap("あいうえおHoldings", 8) == "あいうえお\nHoldings"          # a Latin word is not cut


def test_r5_classic_unit_label_keeps_its_comma():
    from deep_research_tool.chartkit.data_layer import parse_unit
    assert parse_unit("1,000 tons") == ("1,000 tons", 1.0, "")
    assert parse_unit("deaths per 100,000")[0] == "deaths per 100,000"
    assert parse_unit("1,000万円")[:2] == ("円", 1e7)


class TestR5Grounding:

    def test_a_scale_the_page_states_elsewhere_is_not_a_conflict(self):
        for quote, page in [("売上高123,456百万円", "（単位：百万円）\n売上高 118,230 123,456"),
                            ("国内市場規模は1,234億円", "国内市場規模の推移（単位：億円）\n2023年 1,234"),
                            ("revenue of $4,512 million", "(in millions of U.S. dollars) Revenue 4,512 4,876"),
                            ("1,200億円", "2030年に1,200～1,500億円に達する"),
                            ("3.5 million users", "reached 3.5M users"), ("KRW 1.2 trillion", "sales of ₩1.2T")]:
            assert quote_in_source(quote, page), quote

    def test_stated_different_scales_conflict(self):
        for quote, page in [("$5.2 billion", "revenue was $5.2 million"), ("3,000億円", "設備投資額は3,000万円"),
                            ("約3千億円", "国内市場規模は約3千万円と推定される"), ("5千万円", "同社の売上高は5千億円"),
                            ("3百億円", "3百万円"), ("35万人", "利用者は35人")]:
            assert not quote_in_source(quote, page), quote
        assert quote_in_source("約3千億円", "市場は約3千億円")

    def test_half_up_fractions_and_spaced_si_prefix(self):
        from deep_research_tool.chartkit.data_layer import repair_scale
        for v, text in [(13, "8人に1人"), (63, "8分の5"), (3, "40人に1人"), (6.3, "16人に1人"), (12, "8人に1人")]:
            assert value_in_quote(v, "%", text), (v, text)
        assert not value_in_quote(14, "%", "8人に1人") and not value_in_quote(34, "%", "3人に1人")
        for text in ("a 10k ohm pull-up resistor", "10k Ω", "4.7k ohms"):
            value = 4.7 if "4.7" in text else 10
            assert repair_scale(value, text, "kΩ") == ("kΩ", False), text
            assert value_in_quote(value, "kΩ", text), text


class TestR5Facts:

    def test_rate_units_are_not_scaled(self):
        from deep_research_tool.chartkit.facts import _fmt
        assert _fmt(15.7, "10万人当たり") == "15.7（10万人当たり）"
        assert _fmt(15.7, "人口10万人当たり") == "15.7（人口10万人当たり）"
        assert _fmt(5, "10億ドル") == "50億ドル" and _fmt(15.7, "1,000 tons") == "15.7（1,000 tons）"

    def test_topic_scopes_and_rate_topics_are_not_flagged(self):
        bar = panel_facts({"form": "bar", "unit": "億円",
                           "labels": ["toyota", "volkswagen", "hyundai", "honda", "nissan", "bmw"],
                           "series": [{"values": [1030, 900, 720, 400, 330, 250]}]})
        for t in ("欧州メーカーはVolkswagenが首位。", "日系メーカーはToyotaが首位、Nissanが最下位。"):
            assert judgment_conflict(t, bar) is None, t
        assert judgment_conflict("Volkswagenが首位。", bar)
        lv = panel_facts({"form": "line_xy", "unit": "億円", "labels": [], "series": [
            {"values": [{"x": float(2019 + i), "y": v} for i, v in enumerate([1000, 1300, 1100, 1200, 1400])]}]})
        for t in ("増加率は2020年が最大。", "伸びは2020年が最大。", "前年同期比は2020年が最大。", "落ち込みは2021年が最大。"):
            assert judgment_conflict(t, lv) is None, t
        assert judgment_conflict("2020年が過去最高。", lv)


class TestR5Reports:

    @pytest.mark.parametrize("fmt", ["html", "pdf", "markdown"])
    def test_body_link_with_parentheses_passes(self, tmp_path, fmt):
        pytest.importorskip("reportlab")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.semantic_manifest import build_semantic_manifest, check_artifact
        from deep_research_tool.main import _report_titles
        gen, session, locker = TestR3Artifact()._gen(tmp_path)
        body = ("本文 [1]\n\n制度の詳細は[ノンアルコールビール](https://ja.wikipedia.org/wiki/ノンアルコールビール_(日本))"
                "の解説が詳しく、酒税法上の区分と市場の成り立ちを整理している。")
        session.section_contents["1"]["content"] = body
        path = gen.generate_report(session=session, evidence_locker=locker, format=ReportFormat(fmt),
                                   verification_result=None)
        res = check_artifact(build_semantic_manifest({"1": body}), path, locker, extra_allowed=_report_titles(session))
        assert res["status"] == "pass", res
        if fmt == "html":
            assert 'href="https://ja.wikipedia.org/wiki/ノンアルコールビール_(日本)"' in path.read_text(encoding="utf-8")

    def test_list_items_that_start_with_a_number(self):
        from deep_research_tool.report.semantic_manifest import verify_frozen_in_artifact
        text = "2030年までに国内市場は倍増する見込みであり、主要メーカーは増産投資を相次いで進めている。"
        for marker in ("1. ", "1) ", "2. "):
            assert verify_frozen_in_artifact({"chapters": {"1": marker + text}}, text)["ok"], marker

    def test_english_pdf_ignores_unprinted_source_pages(self, tmp_path, monkeypatch):
        pytest.importorskip("reportlab")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.generator import ReportGenerator
        from deep_research_tool.tests.test_finalization_integration import make_locker, make_plan, make_session
        calls = []
        monkeypatch.setattr(ReportGenerator, "_register_japanese_fonts", lambda self: calls.append(1))
        plan = make_plan([("1", "Market share", "d")])
        session = make_session(plan, {"1": {"title": "Market share", "content": "Company A leads. " * 5,
                                            "evidence": []}})
        plan.title = plan.table_of_contents.title = session.query = "US beverage market"
        session.section_contents["1"]["extracted_content"] = [
            {"title": "日本語 | 中文", "content": "言語を選択", "raw_content": "日本市場の資料", "key_points": ["国内"]}]
        gen = ReportGenerator(output_dir=tmp_path, language="en", include_citations=False)
        gen.generate_report(session, make_locker(tmp_path, [("https://ex.com/a", "data " * 20)]),
                            format=ReportFormat.PDF, filename="r")
        assert not calls

    def test_five_digit_values_never_split_in_a_pdf_grid(self, tmp_path):
        pytest.importorskip("reportlab")
        pypdf = pytest.importorskip("pypdf")
        from deep_research_tool.config import ReportFormat
        from deep_research_tool.report.figure_table_generator import FigureTableCollection, TableData
        for n in (9, 11, 12):
            coll = FigureTableCollection(tables=[TableData(
                table_id="w", title="年次推移", caption="", headers=["メーカー"] + [str(2024 - n + i) for i in range(n - 1)],
                rows=[["A社"] + [f"12,{345 + i}" for i in range(n - 1)]], section_id="1")])
            gen, session, locker = TestR3Artifact()._gen(tmp_path / str(n), coll)
            path = gen.generate_report(session=session, evidence_locker=locker, format=ReportFormat.PDF,
                                       verification_result=None)
            text = " ".join("\n".join(p.extract_text() for p in pypdf.PdfReader(str(path)).pages).splitlines())
            assert all(f"12,{345 + i}" in text for i in range(n - 1)), n
