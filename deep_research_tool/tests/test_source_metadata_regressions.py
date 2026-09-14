"""Source provenance and bounded HTTP regressions from the September audit."""

import io
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from unittest.mock import Mock, patch

import pytest

from deep_research_tool.evidence.locker import Evidence, EvidenceLocker, SourceType
from deep_research_tool.evidence.quality_evaluator import QualityEvaluator
from deep_research_tool.evidence.source_metadata import extract_source_metadata
from deep_research_tool.search.base import PageContent
from deep_research_tool.search.duckduckgo import DuckDuckGoSearch


def test_dates_are_distinct_and_arbitrary_body_dates_are_not_publication():
    result = extract_source_metadata("https://example.gov/report", "2035年1月1日には達成予定。公開日: 2025年2月3日。更新日: 2025年4月5日。施行日: 2035年1月1日")
    assert result["published_date"] == "2025-02-03"
    assert result["updated_at"] == "2025-04-05"
    assert result["effective_at"] == "2035-01-01"
    assert result["date_provenance"]["published_date"]["source"] == "body:label"
    assert extract_source_metadata("https://example.gov/report", "2035年1月1日には達成予定")["published_date"] == ""
    assert extract_source_metadata("https://example.gov/report", metadata={"published_date": "2099-01-01"})["published_date"] == ""


def test_explicit_jsonld_and_locker_roundtrip(tmp_path):
    html = '<script type="application/ld+json">{"@type":"Report","datePublished":"2025-04-05","dateModified":"2025-05-06","temporalCoverage":"2024","publisher":{"name":"Office"}}</script>'
    page = PageContent("https://example.gov/report", "Report", "A study of 2024", html_content=html)
    locker = EvidenceLocker(output_dir=tmp_path)
    ev = locker.add_evidence(page.url, page.title, extracted_text=page.text_content, metadata=page.metadata)
    restored = Evidence.from_dict(ev.to_dict())
    assert (restored.published_date, restored.updated_at, restored.data_period) == ("2025-04-05", "2025-05-06", "2024")
    assert restored.publisher == "Office"
    assert restored.source_type == SourceType.OFFICIAL
    assert restored.date_provenance["published_date"]["source"] == "jsonld:datePublished"
    assert page.metadata["retrieved_at"]


def test_duplicate_evidence_preserves_section_associations(tmp_path):
    locker = EvidenceLocker(output_dir=tmp_path)
    a = locker.add_evidence("https://example.gov/report", "Report", "Same text", section_reference="1")
    b = locker.add_evidence("https://example.gov/report", "Report", "Same text", section_reference="2")
    assert a.id == b.id
    assert locker.get_section_evidence("2") == [a]
    assert len(locker.get_all_evidence()) == 1


def test_new_excerpts_share_source_but_new_document_version_does_not(tmp_path):
    locker = EvidenceLocker(output_dir=tmp_path)
    a = locker.add_evidence("https://example.gov/report?utm_source=a", "First title", "First finding", extracted_text="Complete raw text", section_reference="1")
    b = locker.add_evidence("https://example.gov/report", "New title", "Second finding", extracted_text="Complete raw text", section_reference="2")
    c = locker.add_evidence("https://example.gov/report", "New edition", "New finding", extracted_text="Updated raw text", section_reference="3")
    assert a.id == b.id != c.id
    assert len(locker.get_all_evidence()) == 2
    assert "First finding" in a.content_excerpt and "Second finding" in a.content_excerpt
    assert locker.get_section_evidence("2") == [a]


def test_duplicate_dates_never_disagree_with_their_provenance(tmp_path):
    locker = EvidenceLocker(output_dir=tmp_path)
    a = locker.add_evidence("https://example.gov/report", "Report", "Same text", published_date="2020-01-01")
    b = locker.add_evidence("https://example.gov/report", "Report", "Same text", published_date="2025-01-01")
    assert a.id == b.id
    assert b.published_date == "2020-01-01"
    assert b.date_provenance["published_date"]["value"] == b.published_date
    assert b.metadata["provenance_conflicts"][0]["incoming"] == "2025-01-01"


@pytest.mark.parametrize("url", ["https://who.int.example.com/", "https://notreuters.com/", "https://reuters.com.example.org/", "https://example.gov.evil.com/"])
def test_lookalike_domains_not_authoritative(url):
    result = QualityEvaluator().evaluate_url(url)
    assert result["quality_category"].value == "unverified"
    assert result["source_type"].value not in ("official", "news")


def test_real_host_and_subdomains_still_recognized():
    assert QualityEvaluator().evaluate_url("https://www.who.int:443/report")["quality_category"].value == "authoritative"
    assert QualityEvaluator().evaluate_url("https://www.reuters.com/report")["source_type"].value == "news"


def test_fetch_timeout_reaches_transport_and_does_not_change_client_default():
    client = DuckDuckGoSearch(timeout=30, waf_mitigation=False)
    response = Mock(status_code=200)
    with patch("deep_research_tool.search.duckduckgo.requests.get", return_value=response) as get:
        assert client._fetch("https://example.com", {}, timeout=0.5) is response
    passed = get.call_args.kwargs["timeout"]
    assert 0 < passed.total <= 0.5
    assert client.timeout == 30


def test_retry_after_cannot_exceed_fetch_budget():
    client = DuckDuckGoSearch(per_domain_delay=0)
    response = Mock(status_code=429, headers={"Retry-After": "600"})
    session = Mock()
    session.get.return_value = response
    with patch.object(client, "_get_session", return_value=session), patch("deep_research_tool.search.duckduckgo.time.sleep") as sleep:
        with pytest.raises(TimeoutError):
            client._fetch("https://example.com", {}, timeout=0.1)
    assert session.get.call_count == 1
    sleep.assert_not_called()
    response.close.assert_called_once()


def test_trickling_local_response_cannot_extend_body_deadline():
    import threading
    import time
    from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.send_header("Content-Length", "50")
            self.end_headers()
            try:
                for _ in range(50):
                    self.wfile.write(b"x")
                    self.wfile.flush()
                    time.sleep(0.04)
            except OSError:
                pass

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    client = DuckDuckGoSearch(waf_mitigation=False)
    start = time.monotonic()
    try:
        with patch.dict("os.environ", {"NO_PROXY": "127.0.0.1"}), pytest.raises(TimeoutError):
            client._fetch(f"http://127.0.0.1:{server.server_port}/", {}, timeout=0.15)
        assert time.monotonic() - start < 0.7  # full body would take two seconds
    finally:
        server.shutdown()
        server.server_close()


def test_web_docx_preserves_interleaved_blocks():
    from docx import Document
    doc = Document()
    doc.add_paragraph("2024 Revenue")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "100"
    doc.add_paragraph("2025 Revenue")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "200"
    output = io.BytesIO()
    doc.save(output)
    page = DuckDuckGoSearch()._extract_docx_content(Mock(content=output.getvalue()), "https://example.com/report.docx")
    assert page.text_content.index("100") < page.text_content.index("2025 Revenue")
